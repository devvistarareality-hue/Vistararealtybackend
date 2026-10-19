"""Importing investors from Excel: checked row by row, saved as drafts (Approvals → Drafts)."""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from dateutil.relativedelta import relativedelta
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.models import User
from club1000.models import Investor, Scheme
from companies.models import Company


RECENT = date.today() - relativedelta(months=2)          # a term still running — no past terms
RECENT_S, RECENT2_S = RECENT.strftime('%d-%m-%Y'), (date.today() - relativedelta(months=1)).strftime('%d/%m/%Y')


class InvestorImportTests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(code='IMP', name='Imp Co', is_active=True)
        self.admin = User.objects.create_user('a@imp.com', company=self.co, user_code='IMP1', password='p',
                                              name='Admin', role='Admin', modules=['Club 1000'])
        self.scheme = Scheme.objects.create(company=self.co, name='RISE', tenure_months=12, min_ticket_size=100000,
                                            interest_payout_options=['quarterly', 'maturity'],
                                            payout_rates={'quarterly': '15', 'maturity': '15'})
        self.api = APIClient(); self.api.force_authenticate(self.admin)

    def _file(self, rows, headers=None):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Investors'
        ws.append(headers or ['Investor Name', 'Mobile', 'Scheme', 'Amount Invested', 'Investment Date', 'Interest Payout', 'Reference Name', 'Reference Mobile'])
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf)
        return SimpleUploadedFile('inv.xlsx', buf.getvalue())

    def test_template_downloads(self):
        r = self.api.get('/api/club1000/investors/import-template/')
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn('Investors', wb.sheetnames)
        self.assertEqual(wb['Investors'].cell(row=1, column=1).value, 'Investor Name *')   # must-fill marked
        self.assertNotIn('Times Renewed', [c.value for c in wb['Investors'][1]])     # worked out, not typed
        heads = [str(c.value or '').replace(' *', '') for c in wb['Investors'][1]]
        for gone in ('Return %', 'LOI No', 'Security', 'Notes'):     # not asked for in the import
            self.assertNotIn(gone, heads)

    def test_row_missing_a_must_fill_column_is_not_imported(self):
        f = self._file([['Has All', '9800000007', 'RISE', 500000, RECENT_S, '', '', ''],
                        ['No Amount', '9800000006', 'RISE', '', RECENT_S, '', '', '']])
        r = self.api.post('/api/club1000/investors/import/', {'file': f, 'commit': 'true'}, format='multipart')
        self.assertEqual((r.data['created'], r.data['failed']), (1, 1))
        self.assertEqual([i.name for i in Investor.objects.all()], ['Has All'])

    def test_check_reports_errors_and_saves_nothing(self):
        f = self._file([['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', 'Ref', '9811111111'],
                        ['', '98', 'NOPE', 'abc', '2099-01-01', 'weekly', '', '']])
        r = self.api.post('/api/club1000/investors/import/', {'file': f}, format='multipart')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data['ok'], r.data['failed'], r.data['created']), (1, 1, 0))
        bad = r.data['rows'][1]
        self.assertEqual(bad['row'], 3)
        self.assertTrue(any('Scheme' in e for e in bad['errors']))
        self.assertEqual(Investor.objects.count(), 0)

    def test_commit_saves_drafts(self):
        f = self._file([['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', 'Ref', '9811111111'],
                        ['B Two', '9800000002', 'rise', '2,00,000', RECENT2_S, '', '', '']])
        r = self.api.post('/api/club1000/investors/import/', {'file': f, 'commit': 'true'}, format='multipart')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['created'], 2)
        by = {i.name: i for i in Investor.objects.all()}   # names are encrypted — match in Python
        self.assertEqual(by['A One'].total_return_pct, Decimal('15'))   # the scheme's rate for the payout
        self.assertEqual(by['A One'].loi_no or '', '')                  # given when the draft is submitted
        a = by['A One']
        self.assertEqual((a.approval_status, a.interest_payout, a.amount_invested, a.investment_date, a.maturity_date),
                         ('draft', 'quarterly', Decimal('500000'), RECENT, RECENT + relativedelta(months=12)))
        self.assertEqual(a.payouts.count(), 0)          # schedule waits for approval
        b = by['B Two']
        self.assertEqual((b.interest_payout, b.amount_invested), ('quarterly', Decimal('200000')))

    def test_same_investment_twice_is_refused(self):
        row = ['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', '', '']
        self.api.post('/api/club1000/investors/import/', {'file': self._file([row]), 'commit': 'true'}, format='multipart')
        r = self.api.post('/api/club1000/investors/import/', {'file': self._file([row])}, format='multipart')
        self.assertEqual(r.data['failed'], 1)
        self.assertTrue(any('Already in Nexora' in e for e in r.data['rows'][0]['errors']))

    def test_staff_cannot_import(self):
        staff = User.objects.create_user('s@imp.com', company=self.co, user_code='IMP2', password='p',
                                         name='Staff', role='Employee', modules=['Club 1000'])
        api = APIClient(); api.force_authenticate(staff)
        self.assertEqual(api.get('/api/club1000/investors/import-template/').status_code, 403)

    def test_old_investment_is_imported_as_it_is_never_renewed(self):
        """Renewals are not imported (they are done by hand on the website): an old date
        is that investment's own date, and a matured one comes in due for renewal."""
        old = date.today().replace(day=1) - relativedelta(years=6)
        f = self._file([['Ramesh Patel', '9876500000', 'RISE', 5350000, old.strftime('%d-%m-%Y'), 'maturity', '', '']])
        r = self.api.post('/api/club1000/investors/import/', {'file': f, 'commit': 'true'}, format='multipart')
        self.assertEqual((r.data['created'], r.data['failed']), (1, 0), r.data)
        self.assertTrue(any('due for renewal' in w for w in r.data['rows'][0]['warnings']))
        self.assertNotIn('past_terms', r.data['rows'][0]['preview'])
        inv = Investor.objects.get()
        self.assertEqual((inv.renewal_count, inv.past_terms, inv.first_invested_on), (0, None, None))
        self.assertEqual((inv.investment_date, inv.maturity_date), (old, old + relativedelta(months=12)))

    def test_a_past_terms_sheet_in_an_old_file_is_ignored(self):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Investors'
        ws.append(['Investor Name *', 'Mobile *', 'Scheme *', 'Amount Invested *', 'Investment Date *'])
        ws.append(['New One', '9876500001', 'RISE', 500000, RECENT])
        pt = wb.create_sheet('Past Terms')
        pt.append(['Mobile', 'Investor Name', 'Term', 'Term Start', 'Principal (if different)'])
        pt.append(['9876500001', 'New One', None, RECENT - relativedelta(years=1), 10000000])
        buf = io.BytesIO(); wb.save(buf)
        r = self.api.post('/api/club1000/investors/import/', {'file': SimpleUploadedFile('inv.xlsx', buf.getvalue()),
                                                              'commit': 'true'}, format='multipart')
        self.assertEqual((r.data['created'], r.data['failed']), (1, 0), r.data)
        self.assertEqual(Investor.objects.get().renewal_count, 0)

    def test_template_has_no_past_terms_sheet(self):
        r = self.api.get('/api/club1000/investors/import-template/')
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertNotIn('Past Terms', wb.sheetnames)
        self.assertIn('Renewals are not imported', ' '.join(str(c.value or '') for row in wb['Help'].iter_rows() for c in row))

    def test_imported_draft_goes_for_approval_once_its_signed_loi_is_attached(self):
        import base64
        f = self._file([['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', '', '']])
        self.api.post('/api/club1000/investors/import/', {'file': f, 'commit': 'true'}, format='multipart')
        inv = Investor.objects.get()
        drafts = self.api.get('/api/club1000/investors/', {'approval_status': 'draft'}).data
        self.assertEqual([d['id'] for d in drafts], [inv.id])
        self.assertEqual(self.api.get('/api/club1000/investors/', {'approval_status': 'pending'}).data, [])
        self.assertNotEqual(self.api.post(f'/api/club1000/investors/{inv.pk}/action/', {'action': 'approve'},
                                          format='json').status_code, 200)        # a draft cannot be approved
        r = self.api.post('/api/club1000/investors/', {
            'draft_id': inv.id, 'scheme': self.scheme.id, 'name': 'A One', 'phone': '9800000001',
            'amount_invested': '500000', 'investment_date': str(RECENT), 'interest_payout': 'quarterly', 'source': 'walk_in',
            'loi_file': {'name': 'loi.pdf', 'type': 'application/pdf', 'data': base64.b64encode(b'%PDF-1.4 signed').decode()},
        }, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        inv.refresh_from_db()
        self.assertEqual(Investor.objects.count(), 1)                                 # the same row, not a second one
        self.assertEqual(inv.approval_status, 'pending')
        self.assertTrue(inv.loi_document)

    def test_importing_the_same_file_twice_is_caught_while_still_drafts(self):
        row = ['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', '', '']
        self.api.post('/api/club1000/investors/import/', {'file': self._file([row]), 'commit': 'true'}, format='multipart')
        r = self.api.post('/api/club1000/investors/import/', {'file': self._file([row])}, format='multipart')
        self.assertTrue(any('Already in Nexora' in e for e in r.data['rows'][0]['errors']))

    def test_payout_dropdown_follows_the_scheme(self):
        Scheme.objects.create(company=self.co, name='EQUITY', tenure_months=48, min_ticket_size=1,
                              interest_payout_options=['maturity'], payout_rates={'maturity': '100'})
        wb = openpyxl.load_workbook(io.BytesIO(self.api.get('/api/club1000/investors/import-template/').content))
        lists = {r[0]: [v for v in r[6:9] if v] for r in wb['Lists'].iter_rows(values_only=True) if r[0]}
        self.assertEqual(lists, {'EQUITY': ['maturity'], 'RISE': ['quarterly', 'maturity']})
        self.assertEqual([r[1] for r in wb['Lists'].iter_rows(values_only=True) if r[1]], ['quarterly', 'maturity'])   # no monthly
        dv = [d for d in wb['Investors'].data_validations.dataValidation if 'H2' in str(d.sqref)][0]
        self.assertIn('OFFSET(Lists!$G$1', dv.formula1)
        self.assertLess(len(dv.formula1), 255)           # Excel's limit for a list formula

    def test_drafts_can_be_deleted_by_their_owner_and_managers(self):
        f = self._file([['A One', '9800000001', 'RISE', 500000, RECENT_S, 'quarterly', '', ''],
                        ['B Two', '9800000002', 'RISE', 600000, RECENT_S, 'quarterly', '', '']])
        self.api.post('/api/club1000/investors/import/', {'file': f, 'commit': 'true'}, format='multipart')
        ids = list(Investor.objects.values_list('id', flat=True))
        staff = User.objects.create_user('st@imp.com', company=self.co, user_code='IMP3', password='p',
                                         name='Staff', role='Employee', modules=['Club 1000'])
        other = APIClient(); other.force_authenticate(staff)
        # someone else's drafts — refused
        self.assertEqual(other.post('/api/club1000/investors/bulk-delete/', {'ids': ids}, format='json').status_code, 403)
        self.assertEqual(other.delete(f'/api/club1000/investors/{ids[0]}/').status_code, 403)
        # their own draft — allowed, even without manager access
        mine = Investor.objects.create(company=self.co, scheme=self.scheme, added_by=staff, name='Mine', phone='9800000003',
                                       amount_invested=Decimal('500000'), investment_date=RECENT, approval_status='draft')
        self.assertIn(other.delete(f'/api/club1000/investors/{mine.pk}/').status_code, (200, 204))
        # the manager deletes all the imported drafts at once
        r = self.api.post('/api/club1000/investors/bulk-delete/', {'ids': ids}, format='json')
        self.assertEqual((len(r.data['deleted']), len(r.data['refused'])), (2, 0), r.data)
        self.assertEqual(Investor.objects.count(), 0)
        from activity.models import ActivityLog
        self.assertTrue(ActivityLog.objects.filter(summary__startswith='Deleted 2 draft investors').exists()
                        or any(str(x.summary).startswith('Deleted 2 draft investors') for x in ActivityLog.objects.all()))

    def test_only_manager_role_and_above_may_delete_an_investment(self):
        inv = Investor.objects.create(company=self.co, scheme=self.scheme, added_by=self.admin, name='Live One', phone='9800000005',
                                      amount_invested=Decimal('500000'), investment_date=RECENT, approval_status='approved')
        # an Employee given manager-level Club 1000 access (like Shrey Patel) — refused
        emp = User.objects.create_user('e@imp.com', company=self.co, user_code='IMP4', password='p', name='Exec',
                                       role='Employee', modules=['Club 1000'], manager_modules=['Club 1000'])
        api = APIClient(); api.force_authenticate(emp)
        self.assertEqual(api.delete(f'/api/club1000/investors/{inv.pk}/').status_code, 403)
        self.assertEqual(api.post('/api/club1000/investors/bulk-delete/', {'ids': [inv.pk]}, format='json').status_code, 403)
        # a Manager by role — allowed
        mgr = User.objects.create_user('mg@imp.com', company=self.co, user_code='IMP5', password='p', name='Mgr',
                                       role='Manager', modules=['Club 1000'], manager_modules=['Club 1000'])
        api = APIClient(); api.force_authenticate(mgr)
        self.assertIn(api.delete(f'/api/club1000/investors/{inv.pk}/').status_code, (200, 204))
