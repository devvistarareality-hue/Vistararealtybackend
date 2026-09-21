import json
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import is_platform_admin, scope_to_company
from sales.models import Booking
from .engine import rupees, AGEING_BUCKETS
from .models import ARAccount, ARReceipt, ARReceiptAudit
from .permissions import has_ar_access
from .services import (SUSPECT_BELOW, compute_account, expected_collectable, parse_date, pending_revision_ids,
                       sync_accounts, _d)

ZERO = Decimal('0')
MODES = dict(ARReceipt.MODES)


def _deny():
    return Response({'detail': 'You do not have access to Accounts Receivable.'}, status=status.HTTP_403_FORBIDDEN)


def _as_of(request):
    return parse_date(request.query_params.get('as_of')) or timezone.localdate()


def _accounts_qs(request):
    qs = scope_to_company(ARAccount.objects.all(), request.user)
    cid = request.query_params.get('company_id')
    if cid and is_platform_admin(request.user):
        qs = qs.filter(company_id=cid)
    return qs


def _sync(request):
    qs = scope_to_company(Booking.objects.all(), request.user)
    cid = request.query_params.get('company_id')
    if cid and is_platform_admin(request.user):
        qs = qs.filter(company_id=cid)
    sync_accounts(qs.filter(status='sold'))


def _summary(acct, plan, r, mismatch):
    b = acct.booking
    no_schedule = not any(l.kind in ('inst', 'nsd', 'extra') for l in plan)
    return {
        'id': acct.id,
        'status': acct.status,
        'booking_id': b.id,
        'project_id': b.project_id,
        'project': b.project.name if b.project_id else '',
        'plots': b.plot_numbers or (b.plot.number if b.plot_id else ''),
        'client_name': b.client_name or '',
        'phone': b.phone or '',
        'booking_date': b.booking_date.isoformat() if b.booking_date else None,
        'total_deal': rupees(_d(b.final_amount)),
        'stamp_duty': rupees(_d(b.stamp_duty)),
        'reg_fees': rupees(_d(b.reg_fees)),
        'collectable': rupees(r.collectable),
        'received': rupees(r.received),
        'pct_realised': round(float(r.received / r.collectable * 100), 2) if r.collectable else 0,
        'outstanding': rupees(r.outstanding),
        'overdue': rupees(r.overdue),
        'not_due': rupees(r.not_due),
        'net_interest': rupees(r.net_interest),
        'os_with_interest': rupees(r.os_with_interest),
        'ageing': {k: rupees(v) for k, v in r.ageing.items()},
        'legal_due_date': acct.legal_due_date.isoformat() if acct.legal_due_date else None,
        # The plan should add up to Total Deal − Stamp − Reg; flag it when it doesn't
        # so Accounts can see the LOI schedule doesn't cover the whole deal.
        # Installment amounts are typed by hand, so a few rupees of rounding is normal;
        # only a real gap (over ₹10) is worth a warning.
        'plan_mismatch': 0 if no_schedule else (rupees(mismatch) if abs(mismatch) > 10 else 0),
        # No installments on the booking (an EOI, or an LOI whose schedule was never
        # entered): AR can't track dues until Sales enters one, so say that plainly
        # instead of showing a huge "plan mismatch".
        'no_schedule': no_schedule,
        # A deal under ₹1 lakh is a mistyped amount, not a real sale.
        'suspect_amount': expected_collectable(b) < SUSPECT_BELOW,
        # A newer revision of this deal is awaiting approval; AR shows the last approved one.
        'revision_pending': bool(getattr(acct, '_rev_pending', False)),
    }


# Just the booking columns the register and ledger use. A Booking row carries ~30
# encrypted fields and each one is decrypted on load; the register loads every
# account, so loading only these is most of its speed.
_BOOKING_COLS = (
    'booking__id', 'booking__company_id', 'booking__project_id', 'booking__plot_id', 'booking__plot_numbers',
    'booking__client_name', 'booking__phone', 'booking__booking_date', 'booking__final_amount',
    'booking__stamp_duty', 'booking__reg_fees', 'booking__total_extra', 'booking__installments',
    'booking__extra_work_inst', 'booking__cancelled_at', 'booking__project__name', 'booking__plot__number',
)
_ACCOUNT_COLS = ('id', 'company_id', 'root_booking_id', 'booking_id', 'status', 'frozen_at', 'legal_due_date')


def _slim(qs):
    return qs.select_related('booking', 'booking__project', 'booking__plot').only(*_ACCOUNT_COLS, *_BOOKING_COLS)


def _computed(qs, as_of):
    accts = list(_slim(qs).prefetch_related('receipts'))
    pending = pending_revision_ids(a.booking_id for a in accts)
    out = []
    for acct in accts:
        acct._rev_pending = acct.booking_id in pending
        receipts = [x for x in acct.receipts.all() if not x.is_deleted]
        plan, r, mismatch = compute_account(acct, as_of, receipts)
        out.append((acct, plan, r, mismatch, receipts))
    return out


class ARRegisterView(APIView):
    """The Plot Master equivalent: one row per account."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not has_ar_access(request.user):
            return _deny()
        _sync(request)
        as_of = _as_of(request)
        qs = _accounts_qs(request)
        pid = request.query_params.get('project')
        if pid:
            qs = qs.filter(booking__project_id=pid)
        rows = [_summary(a, p, r, m) for a, p, r, m, _ in _computed(qs, as_of)]
        q = (request.query_params.get('q') or '').strip().lower()
        if q:
            rows = [x for x in rows if q in x['client_name'].lower() or q in x['phone'] or q in str(x['plots']).lower()]
        if request.query_params.get('overdue') == '1':
            rows = [x for x in rows if x['overdue'] > 0]
        rows.sort(key=lambda x: (x['project'], str(x['plots'])))
        return Response({'as_of': as_of.isoformat(), 'results': rows})


class ARAccountView(APIView):
    """The Ledger Format equivalent for one account."""
    permission_classes = [IsAuthenticated]

    def _get(self, request, pk):
        return _slim(_accounts_qs(request)).filter(pk=pk).first()

    def get(self, request, pk):
        if not has_ar_access(request.user):
            return _deny()
        acct = self._get(request, pk)
        if not acct:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        as_of = _as_of(request)
        receipts = sorted(acct.receipts.filter(is_deleted=False).select_related('created_by'), key=lambda x: (x.paid_on, x.id))
        plan, r, mismatch = compute_account(acct, as_of, receipts)
        acct._rev_pending = acct.booking_id in pending_revision_ids([acct.booking_id])
        data = _summary(acct, plan, r, mismatch)
        data.update({
            'as_of': as_of.isoformat(),
            'plan': [{
                'key': s.line.key, 'no': s.line.no, 'label': s.line.label, 'kind': s.line.kind,
                'due': s.line.due.isoformat() if s.line.due else None,
                'amount': rupees(s.line.amount), 'paid': rupees(s.paid),
                'pending': rupees(s.remaining), 'status': s.status,
            } for s in r.lines],
            'receipts': [{
                'id': x.id, 'paid_on': x.paid_on.isoformat(), 'amount': float(_d(x.amount)),
                'mode': x.mode, 'mode_label': MODES.get(x.mode, x.mode), 'remarks': x.remarks or '',
                'source': x.source, 'created_by': x.created_by.name if x.created_by_id else '',
            } for x in receipts],
            'interest_rows': [{
                'inst_no': row.inst_no, 'due': row.due.isoformat() if row.due else None,
                'paid_on': row.paid_on.isoformat() if row.paid_on else None,
                'amount': rupees(row.amount), 'days': row.days, 'interest': rupees(row.interest),
            } for row in r.rows],
            'overpaid': rupees(r.overpaid),
            'overpaid_credit': rupees(r.overpaid_credit),
            'month_forecast': [{'label': lbl, 'amount': rupees(v)} for lbl, v in r.month_forecast],
        })
        return Response(data)

    def patch(self, request, pk):
        """Set (or clear) the due date of the Legal & Other Charges line."""
        if not has_ar_access(request.user):
            return _deny()
        acct = self._get(request, pk)
        if not acct:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        if 'legal_due_date' in request.data:
            raw = request.data.get('legal_due_date')
            d = parse_date(raw) if raw else None
            if raw and d is None:
                return Response({'detail': 'Invalid date.'}, status=status.HTTP_400_BAD_REQUEST)
            acct.legal_due_date = d
            acct.save(update_fields=['legal_due_date', 'updated_at'])
        return self.get(request, pk)


def _snap(rc):
    return json.dumps({'paid_on': rc.paid_on.isoformat() if rc.paid_on else None,
                       'amount': str(_d(rc.amount)), 'mode': rc.mode, 'remarks': rc.remarks or ''})


def _clean_receipt(data, partial=False):
    out, errs = {}, {}
    if 'paid_on' in data or not partial:
        d = parse_date(data.get('paid_on'))
        if not d:
            errs['paid_on'] = 'A valid payment date is required.'
        elif d > timezone.localdate():
            errs['paid_on'] = 'Payment date cannot be in the future.'
        else:
            out['paid_on'] = d
    if 'amount' in data or not partial:
        a = _d(data.get('amount'))
        if a <= 0:
            errs['amount'] = 'Amount must be greater than zero.'
        else:
            out['amount'] = a
    if 'mode' in data or not partial:
        m = (data.get('mode') or '').lower()
        if m not in MODES:
            errs['mode'] = 'Mode must be Bank, NBFC, Cash or Cheque.'
        else:
            out['mode'] = m
    if 'remarks' in data:
        out['remarks'] = (data.get('remarks') or '').strip()[:500]
    return out, errs


class ARReceiptCreateView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        if not has_ar_access(request.user):
            return _deny()
        acct = _accounts_qs(request).filter(pk=pk).first()
        if not acct:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        if acct.status == 'frozen':
            return Response({'detail': 'This booking was cancelled — its account is frozen.'}, status=status.HTTP_400_BAD_REQUEST)
        vals, errs = _clean_receipt(request.data)
        if errs:
            return Response(errs, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            rc = ARReceipt.objects.create(account=acct, created_by=request.user, updated_by=request.user, **vals)
            ARReceiptAudit.objects.create(receipt=rc, action='create', changed_by=request.user, after=_snap(rc))
        return Response({'id': rc.id}, status=status.HTTP_201_CREATED)


class ARReceiptView(APIView):
    permission_classes = [IsAuthenticated]

    def _get(self, request, rid):
        accts = _accounts_qs(request).values('id')
        return ARReceipt.objects.filter(pk=rid, account_id__in=accts, is_deleted=False).select_related('account').first()

    def patch(self, request, rid):
        if not has_ar_access(request.user):
            return _deny()
        rc = self._get(request, rid)
        if not rc:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        vals, errs = _clean_receipt(request.data, partial=True)
        if errs:
            return Response(errs, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            before = _snap(rc)
            for k, v in vals.items():
                setattr(rc, k, v)
            rc.updated_by = request.user
            rc.save()
            ARReceiptAudit.objects.create(receipt=rc, action='update', changed_by=request.user, before=before, after=_snap(rc))
        return Response({'id': rc.id})

    def delete(self, request, rid):
        if not has_ar_access(request.user):
            return _deny()
        rc = self._get(request, rid)
        if not rc:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        with transaction.atomic():
            before = _snap(rc)
            rc.is_deleted = True
            rc.updated_by = request.user
            rc.save(update_fields=['is_deleted', 'updated_by', 'updated_at'])
            ARReceiptAudit.objects.create(receipt=rc, action='delete', changed_by=request.user, before=before)
        return Response(status=status.HTTP_204_NO_CONTENT)


class ARReceiptAuditView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, rid):
        if not has_ar_access(request.user):
            return _deny()
        accts = _accounts_qs(request).values('id')
        rc = ARReceipt.objects.filter(pk=rid, account_id__in=accts).first()
        if not rc:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        return Response([{
            'action': a.action, 'changed_by': a.changed_by.name if a.changed_by_id else '',
            'changed_at': a.changed_at.isoformat(),
            'before': json.loads(a.before) if a.before else None,
            'after': json.loads(a.after) if a.after else None,
        } for a in rc.audit.select_related('changed_by')])


class ARDashboardView(APIView):
    """Portfolio view: totals, ageing, month-wise dues, worst accounts and the
    accounts whose booking data needs fixing."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not has_ar_access(request.user):
            return _deny()
        _sync(request)
        as_of = _as_of(request)
        base = _accounts_qs(request).filter(status='active')
        projects = {}
        for pid, name in base.values_list('booking__project_id', 'booking__project__name').distinct():
            if pid:
                projects[pid] = name or ''
        qs = base
        pid = request.query_params.get('project')
        if pid:
            qs = qs.filter(booking__project_id=pid)
        totals = {k: ZERO for k in ('collectable', 'received', 'outstanding', 'overdue', 'not_due', 'net_interest', 'os_with_interest')}
        ageing = {label: ZERO for label, _, _ in AGEING_BUCKETS}
        forecast, order, rows = {}, [], []
        issues = {'no_schedule': 0, 'plan_mismatch': 0, 'suspect_amount': 0, 'revision_pending': 0}
        for acct, plan, r, m, _ in _computed(qs, as_of):
            for k in totals:
                totals[k] += getattr(r, k)
            for k, v in r.ageing.items():
                ageing[k] += v
            for lbl, v in r.month_forecast:
                if lbl not in forecast:
                    order.append(lbl)
                    forecast[lbl] = ZERO
                forecast[lbl] += v
            sm = _summary(acct, plan, r, m)
            issues['no_schedule'] += sm['no_schedule']
            issues['plan_mismatch'] += bool(sm['plan_mismatch'])
            issues['suspect_amount'] += sm['suspect_amount']
            issues['revision_pending'] += sm['revision_pending']
            rows.append((sm, r.ageing.get('>180', ZERO)))

        def brief(sm, amount):
            return {'id': sm['id'], 'client': sm['client_name'] or '—', 'project': sm['project'],
                    'plots': sm['plots'], 'amount': rupees(amount)}

        top_overdue = sorted((x for x in rows if x[0]['overdue'] > 0), key=lambda x: x[0]['overdue'], reverse=True)[:10]
        top_180 = sorted((x for x in rows if x[1] > 0), key=lambda x: x[1], reverse=True)[:10]
        n = len(rows)
        return Response({
            'as_of': as_of.isoformat(),
            'accounts': n,
            'overdue_accounts': sum(1 for x in rows if x[0]['overdue'] > 0),
            'projects': [{'id': k, 'name': v} for k, v in sorted(projects.items(), key=lambda kv: kv[1])],
            'totals': {k: rupees(v) for k, v in totals.items()},
            'pct_realised': round(float(totals['received'] / totals['collectable'] * 100), 1) if totals['collectable'] else 0,
            'ageing': {k: rupees(v) for k, v in ageing.items()},
            'month_forecast': [{'label': lbl, 'amount': rupees(forecast[lbl])} for lbl in order],
            'top_overdue': [brief(sm, sm['overdue']) for sm, _ in top_overdue],
            'top_over_180': [brief(sm, v) for sm, v in top_180],
            'issues': issues,
        })


class ARStatementView(APIView):
    """The account statement a client can be sent, as a print-ready HTML page.
    The web prints it to PDF in the browser; the app turns it into a PDF with
    expo-print. One template, so both look the same."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        if not has_ar_access(request.user):
            return _deny()
        acct = _slim(_accounts_qs(request)).select_related('company').filter(pk=pk).first()
        if not acct:
            return Response({'detail': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
        as_of = _as_of(request)
        receipts = sorted(acct.receipts.filter(is_deleted=False), key=lambda x: (x.paid_on, x.id))
        plan, r, mismatch = compute_account(acct, as_of, receipts)
        sm = _summary(acct, plan, r, mismatch)
        fmt = lambda d: d.strftime('%d/%m/%Y') if d else '—'
        html = render_to_string('receivables/statement.html', {
            'company': acct.company, 's': sm, 'as_of': fmt(as_of), 'booked': fmt(acct.booking.booking_date),
            'pct_bar': max(0, min(100, sm['pct_realised'])), 'generated': timezone.localtime().strftime('%d/%m/%Y %I:%M %p'),
            'plan': [{'no': x.line.no, 'label': x.line.label, 'due': fmt(x.line.due), 'amount': rupees(x.line.amount),
                      'paid': rupees(x.paid), 'pending': rupees(x.remaining), 'status': x.status} for x in r.lines],
            # Remarks are internal notes ("collected by …"), not for the client.
            'receipts': [{'date': fmt(x.paid_on), 'amount': rupees(_d(x.amount)), 'mode': MODES.get(x.mode, x.mode)}
                         for x in receipts],
            'rows': [{'inst': x.inst_no, 'due': fmt(x.due), 'paid': fmt(x.paid_on) if x.paid_on else 'Unpaid',
                      'amount': rupees(x.amount), 'days': '—' if x.days is None else x.days, 'interest': rupees(x.interest)} for x in r.rows],
            'overpaid': rupees(r.overpaid), 'overpaid_credit': rupees(r.overpaid_credit),
        })
        return HttpResponse(html, content_type='text/html; charset=utf-8')


# ── Excel import: past receipts only ─────────────────────────────────────────
# The old workbooks' Payment Tracker has one row per allocation, not per payment:
# a receipt that spilled across installments appears once with its Paid amount and
# again as Carry-In rows with Paid blank. So only rows with Paid > 0 are real
# receipts — which is also all the user asked to import; the plan comes from the LOI.
HEADER_ALIASES = {
    'plot': ('plot no', 'plot', 'plot number', 'unit', 'unit no'),
    'paid_on': ('paid dates', 'paid date', 'payment date', 'date'),
    'amount': ('paid', 'amount', 'paid amount'),
    'mode': ('mode', 'payment mode'),
    'remarks': ('remarks', 'remark', 'narration'),
}
MODE_ALIASES = {'bank': 'bank', 'nbfc': 'nbfc', 'cash': 'cash', 'cheque': 'cheque', 'chq': 'cheque', 'check': 'cheque'}


def _norm_plot(v):
    s = str(v or '').strip().upper()
    if s.endswith('.0'):
        s = s[:-2]
    return s


def _plots_of(b):
    raw = b.plot_numbers or (b.plot.number if b.plot_id and b.plot else '')
    out = set()
    for p in str(raw).split(','):
        p = _norm_plot(p)
        if p:
            out.add(p)
            # "A-57" should also match a sheet that just says 57.
            tail = p.split('-')[-1]
            out.add(tail)
    return out


def _read_rows(fileobj):
    import openpyxl
    wb = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == 'payment tracker'), wb[wb.sheetnames[0]])
    rows = ws.iter_rows(values_only=True)
    header_idx, cols = None, {}
    for i, row in enumerate(rows):
        names = [str(c or '').strip().lower() for c in row]
        found = {}
        for key, aliases in HEADER_ALIASES.items():
            for j, n in enumerate(names):
                if n in aliases:
                    found[key] = j
                    break
        if {'plot', 'paid_on', 'amount'} <= set(found):
            header_idx, cols = i, found
            break
        if i > 20:
            break
    if header_idx is None:
        return None, 'Could not find the header row (need Plot No, Paid Date/Paid Dates and Paid/Amount columns).'
    out = []
    for n, row in enumerate(rows, start=header_idx + 2):
        def cell(k):
            j = cols.get(k)
            return row[j] if j is not None and j < len(row) else None
        amount = _d(cell('amount'))
        if amount <= 0:
            continue  # carry-in / budget rows, blank rows
        out.append({
            'line': n, 'plot': _norm_plot(cell('plot')), 'paid_on': parse_date(cell('paid_on')),
            'amount': amount, 'mode': MODE_ALIASES.get(str(cell('mode') or '').strip().lower(), ''),
            'remarks': str(cell('remarks') or '').strip()[:500],
        })
    return out, None


def _check_rows(request, pid, rows):
    """Validate receipt rows (from an uploaded Excel file) for one project. Each
    row names its plot. Returns (ok, skipped); ok rows carry
    account_id. Same plot + date + amount as a receipt already on the account is
    a duplicate, so re-running an import adds nothing."""
    _sync(request)
    accts = list(_slim(_accounts_qs(request).filter(booking__project_id=pid)))
    by_id = {a.id: a for a in accts}
    by_plot = {}
    for a in accts:
        for p in _plots_of(a.booking):
            by_plot.setdefault(p, a)
    existing = {(r.account_id, r.paid_on, _d(r.amount)) for r in ARReceipt.objects.filter(account__in=accts, is_deleted=False)}
    today = timezone.localdate()
    ok, skipped = [], []
    for row in rows:
        acct = by_id.get(row.get('account_id')) if row.get('account_id') else by_plot.get(row.get('plot'))
        reason = None
        if not row['paid_on']:
            reason = 'No valid payment date'
        elif row['paid_on'] > today:
            reason = 'Payment date is in the future'
        elif row['amount'] <= 0:
            reason = 'Amount must be more than zero'
        elif not row['mode']:
            reason = 'Mode must be Bank, NBFC, Cash or Cheque'
        elif not acct:
            reason = f"No approved booking for plot {row.get('plot') or '?'} in this project"
        elif acct.status == 'frozen':
            reason = 'Booking cancelled — account frozen'
        elif (acct.id, row['paid_on'], row['amount']) in existing:
            reason = 'Already recorded (same plot, date and amount)'
        entry = {**row, 'paid_on': row['paid_on'].isoformat() if row['paid_on'] else None,
                 'amount': float(row['amount']), 'account_id': acct.id if acct else None,
                 'plot': row.get('plot') or (str(acct.booking.plot_numbers or '') if acct else ''),
                 'client_name': acct.booking.client_name if acct else ''}
        if reason:
            skipped.append({**entry, 'reason': reason})
        else:
            ok.append(entry)
            existing.add((acct.id, row['paid_on'], row['amount']))
    return ok, skipped


def _save_rows(request, ok, source):
    with transaction.atomic():
        for e in ok:
            rc = ARReceipt.objects.create(
                account_id=e['account_id'], paid_on=parse_date(e['paid_on']), amount=_d(e['amount']),
                mode=e['mode'], remarks=e['remarks'], source=source,
                created_by=request.user, updated_by=request.user)
            ARReceiptAudit.objects.create(receipt=rc, action='create', changed_by=request.user, after=_snap(rc))


def _result(ok, skipped, committed):
    return Response({
        'committed': committed,
        'ready': len(ok), 'skipped': len(skipped),
        'total_amount': rupees(sum((_d(e['amount']) for e in ok), ZERO)),
        'rows': ok, 'skipped_rows': skipped,
    })


def _wants_commit(request):
    return request.data.get('commit') in ('1', 'true', True, 1)


class ARImportView(APIView):
    """POST multipart: file (.xlsx), project_id, commit=1 to save (default preview)."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if not has_ar_access(request.user):
            return _deny()
        f = request.FILES.get('file')
        pid = request.data.get('project_id')
        if not f or not pid:
            return Response({'detail': 'Choose the project and an .xlsx file.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            rows, err = _read_rows(f)
        except Exception:
            return Response({'detail': 'That file could not be read as an Excel workbook (.xlsx).'}, status=status.HTTP_400_BAD_REQUEST)
        if err:
            return Response({'detail': err}, status=status.HTTP_400_BAD_REQUEST)
        ok, skipped = _check_rows(request, pid, rows)
        commit = _wants_commit(request) and bool(ok)
        if commit:
            _save_rows(request, ok, 'import')
        return _result(ok, skipped, commit)


class ARImportTemplateView(APIView):
    """GET ?project_id= → an .xlsx with one row per approved plot of the project,
    ready to fill in and upload on the Import receipts page."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not has_ar_access(request.user):
            return _deny()
        pid = request.query_params.get('project_id')
        if not pid:
            return Response({'detail': 'Choose a project.'}, status=status.HTTP_400_BAD_REQUEST)
        import io
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        _sync(request)
        accts = [a for a in _slim(_accounts_qs(request).filter(booking__project_id=pid, status='active'))]
        accts.sort(key=lambda a: _plot_sort_key(str(a.booking.plot_numbers or '')))
        project = accts[0].booking.project.name if accts else 'Project'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Payment Tracker'
        head = ['Plot No', 'Client', 'Paid Date', 'Paid', 'Mode', 'Remarks']
        ws.append(head)
        fill = PatternFill('solid', fgColor='1F5490')
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = fill
            c.alignment = Alignment(vertical='center')
        for a in accts:
            ws.append([str(a.booking.plot_numbers or ''), a.booking.client_name or '', None, None, None, None])
        for col, w in zip('ABCDEF', (12, 32, 14, 16, 12, 40)):
            ws.column_dimensions[col].width = w
        n = ws.max_row
        for r in range(2, max(n, 2) + 200):
            ws.cell(r, 3).number_format = 'DD/MM/YYYY'
            ws.cell(r, 4).number_format = '#,##,##0'
        dv = DataValidation(type='list', formula1='"Bank,NBFC,Cash,Cheque"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'E2:E{max(n, 2) + 200}')
        ws.freeze_panes = 'A2'
        note = wb.create_sheet('How to fill')
        for line in (
            f'{project} — receipts template',
            'One row per payment. For a second payment on the same plot, add a row with the same Plot No.',
            'Paid Date: the date the money was received (not in the future).',
            'Paid: the amount in rupees. Mode: Bank, NBFC, Cash or Cheque.',
            'Leave a plot row blank if nothing was received — blank rows are ignored.',
            'Upload this file on the Import page; you will see a preview before anything is saved.',
        ):
            note.append([line])
        note.column_dimensions['A'].width = 100
        buf = io.BytesIO()
        wb.save(buf)
        name = f"AR receipts - {project}.xlsx".replace('/', '-')
        resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{name}"'
        return resp


def _plot_sort_key(p):
    import re
    m = re.match(r'([A-Za-z-]*)(\d+)', p.strip())
    return (m.group(1).upper(), int(m.group(2))) if m else (p.upper(), 0)
