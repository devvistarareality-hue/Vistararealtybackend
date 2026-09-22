"""AR API tests — run against local SQLite, never the production database:

    DATABASE_URL= DB_ENGINE=django.db.backends.sqlite3 ./venv/bin/python manage.py test receivables
"""
import io
from datetime import date, timedelta
from decimal import Decimal as D

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import User
from companies.models import Company
from sales.models import Booking, Project
from receivables.models import ARAccount, ARReceipt, ARReceiptAudit


def make_booking(company, project, plot='10', **kw):
    base = dict(
        company=company, project=project, plot_numbers=plot, client_name='Jigar Makwana', phone='8160191851',
        status='sold', accounts_status='approved',
        installments=[
            {'no': 1, 'date': '2025-08-01', 'amt': 100000},
            {'no': 2, 'date': '2025-09-12', 'amt': 4050000},
            {'no': 3, 'date': '2025-09-30', 'amt': 2500000},
            {'no': 4, 'date': '2025-10-12', 'amt': 4875000},
            {'no': 5, 'date': '2025-11-12', 'amt': 4875000},
        ],
        # Legal & Other = total_extra − stamp − reg = 3,16,923 (Legal 40,000 + Maint 2,76,923)
        total_extra=D('316923'), stamp_duty=D('0'), reg_fees=D('0'), final_amount=D('16716923'),
    )
    base.update(kw)
    return Booking.objects.create(**base)


class ARApiTests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(code='VIS', name='Vistara')
        self.other = Company.objects.create(code='OTH', name='Other')
        self.project = Project.objects.create(company=self.co, name='Kalrav 2')
        self.user = User.objects.create_user('ar1@test.local', company=self.co, user_code='AR1', password='x', name='AR User',
                                             role='Employee', modules=['AR'])
        self.nobody = User.objects.create_user('no1@test.local', company=self.co, user_code='NO1', password='x', name='No AR',
                                               role='Employee', modules=['Sales'])
        self.booking = make_booking(self.co, self.project)
        self.api = APIClient()
        self.api.force_authenticate(self.user)

    def acct(self):
        rows = self.api.get('/api/ar/accounts/').json()['results']
        self.assertEqual(len(rows), 1)
        return rows[0]

    def test_approved_booking_becomes_an_account_with_its_plan(self):
        a = self.acct()
        self.assertEqual(a['collectable'], 16716923)
        self.assertEqual(a['plan_mismatch'], 0)
        ledger = self.api.get(f"/api/ar/accounts/{a['id']}/").json()
        self.assertEqual([p['no'] for p in ledger['plan']], ['1', '2', '3', '4', '5', 'L'])
        self.assertIsNone(ledger['plan'][-1]['due'])  # legal line has no date until set

    def test_not_yet_accounts_approved_is_excluded(self):
        self.booking.accounts_status = 'pending'
        self.booking.save()
        self.assertEqual(self.api.get('/api/ar/accounts/').json()['results'], [])

    def test_access_is_module_gated_and_company_scoped(self):
        c = APIClient()
        c.force_authenticate(self.nobody)
        self.assertEqual(c.get('/api/ar/accounts/').status_code, 403)
        a = self.acct()
        outsider = User.objects.create_user('ot1@test.local', company=self.other, user_code='OT1', password='x', name='O',
                                            role='Employee', modules=['AR'])
        c.force_authenticate(outsider)
        self.assertEqual(c.get(f"/api/ar/accounts/{a['id']}/").status_code, 404)

    def test_receipt_lifecycle_is_audited(self):
        a = self.acct()
        r = self.api.post(f"/api/ar/accounts/{a['id']}/receipts/",
                          {'paid_on': '2025-08-01', 'amount': '100000', 'mode': 'nbfc', 'remarks': 'first'}, format='json')
        self.assertEqual(r.status_code, 201)
        rid = r.json()['id']
        self.assertEqual(self.api.get(f"/api/ar/accounts/{a['id']}/").json()['received'], 100000)

        self.api.patch(f'/api/ar/receipts/{rid}/', {'amount': '90000'}, format='json')
        self.api.delete(f'/api/ar/receipts/{rid}/')
        self.assertEqual(self.api.get(f"/api/ar/accounts/{a['id']}/").json()['received'], 0)
        self.assertTrue(ARReceipt.objects.get(pk=rid).is_deleted)  # soft delete
        audit = self.api.get(f'/api/ar/receipts/{rid}/audit/').json()
        self.assertEqual([x['action'] for x in audit], ['delete', 'update', 'create'])
        self.assertEqual(audit[1]['before']['amount'], '100000')
        self.assertEqual(audit[1]['after']['amount'], '90000')

    def test_receipt_validation(self):
        a = self.acct()
        future = (timezone.localdate() + timedelta(days=2)).isoformat()
        r = self.api.post(f"/api/ar/accounts/{a['id']}/receipts/", {'paid_on': future, 'amount': '0', 'mode': 'upi'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(set(r.json()), {'paid_on', 'amount', 'mode'})

    def test_cancellation_freezes_the_account(self):
        a = self.acct()
        ARReceipt.objects.create(account_id=a['id'], paid_on=date(2025, 8, 1), amount=D('100000'), mode='bank')
        self.booking.cancelled_at = timezone.now()
        self.booking.save()
        # No longer approved, so it leaves AR (like the Bookings page)…
        self.assertEqual(self.api.get('/api/ar/accounts/').json()['results'], [])
        r = self.api.post(f"/api/ar/accounts/{a['id']}/receipts/", {'paid_on': '2025-09-01', 'amount': '5', 'mode': 'bank'}, format='json')
        self.assertEqual(r.status_code, 404)
        # …but the account is frozen, not deleted, and its receipts are kept.
        acct = ARAccount.objects.get(pk=a['id'])
        self.assertEqual(acct.status, 'frozen')
        self.assertEqual(acct.receipts.filter(is_deleted=False).count(), 1)

    def test_revision_replaces_plan_receipts_stay(self):
        a = self.acct()
        ARReceipt.objects.create(account_id=a['id'], paid_on=date(2025, 8, 1), amount=D('100000'), mode='bank')
        rev = make_booking(self.co, self.project, revision_of=self.booking, revision_no=1,
                           installments=[{'no': 1, 'date': '2025-08-01', 'amt': 16400000}])
        a2 = self.acct()
        self.assertEqual(a2['id'], a['id'])
        self.assertEqual(a2['booking_id'], rev.id)
        self.assertEqual(a2['received'], 100000)
        ledger = self.api.get(f"/api/ar/accounts/{a['id']}/").json()
        self.assertEqual([p['no'] for p in ledger['plan']], ['1', 'L'])

    def test_set_legal_due_date(self):
        a = self.acct()
        r = self.api.patch(f"/api/ar/accounts/{a['id']}/", {'legal_due_date': '2025-12-12'}, format='json').json()
        self.assertEqual(r['plan'][-1]['due'], '2025-12-12')

    def test_plot10_through_the_api(self):
        """The workbook's Plot 10, end to end: receipts in, ledger out."""
        a = self.acct()
        self.api.patch(f"/api/ar/accounts/{a['id']}/", {'legal_due_date': '2025-12-12'}, format='json')
        for d, amt, mode in [('2025-08-01', 100000, 'nbfc'), ('2025-09-12', 4050000, 'nbfc'), ('2025-09-28', 3000000, 'nbfc'),
                             ('2025-11-25', 2000000, 'nbfc'), ('2025-12-20', 1500000, 'nbfc'), ('2025-12-26', 2000000, 'bank'),
                             ('2026-01-03', 500000, 'bank'), ('2026-02-01', 1000000, 'nbfc'), ('2026-03-31', 1000000, 'nbfc'),
                             ('2026-07-16', 1567000, 'nbfc')]:
            ARReceipt.objects.create(account_id=a['id'], paid_on=date.fromisoformat(d), amount=D(amt), mode=mode)
        led = self.api.get(f"/api/ar/accounts/{a['id']}/?as_of=2026-09-21").json()
        self.assertEqual((led['received'], led['outstanding'], led['net_interest'], led['os_with_interest']),
                         (16717000, -77, 615052, 614975))

    def _tracker_xlsx(self):
        """Mimics the old workbook's Payment Tracker: one row per allocation, with
        Carry-In rows whose Paid is blank — those must NOT become receipts."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Payment Tracker'
        ws.append(['Project', 'Plot No', 'Installment No', 'Due Date', 'Budgeted', 'Carry-In', 'Net Due',
                   'Paid', 'Paid Dates', 'Mode', 'Remarks'])
        ws.append(['K2', 10, 1, '01/08/25', 100000, None, 100000, 100000, date(2025, 8, 1), 'NBFC', ''])
        ws.append(['K2', 10, 1, '01/08/25', None, None, None, 4050000, date(2025, 9, 12), 'NBFC', ''])
        ws.append(['K2', 10, 2, '12/09/25', 4050000, 4050000, None, None, date(2025, 9, 12), 'NBFC', '-'])  # carry-in
        ws.append(['K2', 10, 4, '12/10/25', 875000, None, 875000, 2000000, date(2025, 12, 26), 'BANK', 'TD-CHQ'])
        ws.append(['K2', 99, 1, '01/08/25', 1, None, 1, 50000, date(2025, 8, 1), 'BANK', 'unknown plot'])
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile('ar.xlsx', buf.getvalue(),
                                  content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_import_takes_only_real_receipts_and_is_idempotent(self):
        self.acct()
        preview = self.api.post('/api/ar/import/', {'file': self._tracker_xlsx(), 'project_id': self.project.id}, format='multipart').json()
        self.assertFalse(preview['committed'])
        self.assertEqual(preview['ready'], 3)                     # carry-in row ignored
        self.assertEqual(preview['total_amount'], 6150000)
        self.assertEqual([s['plot'] for s in preview['skipped_rows']], ['99'])
        self.assertEqual(ARReceipt.objects.count(), 0)            # preview writes nothing

        done = self.api.post('/api/ar/import/', {'file': self._tracker_xlsx(), 'project_id': self.project.id, 'commit': '1'}, format='multipart').json()
        self.assertTrue(done['committed'])
        self.assertEqual(ARReceipt.objects.filter(source='import').count(), 3)
        self.assertEqual(ARReceiptAudit.objects.filter(action='create').count(), 3)

        again = self.api.post('/api/ar/import/', {'file': self._tracker_xlsx(), 'project_id': self.project.id, 'commit': '1'}, format='multipart').json()
        self.assertEqual(again['ready'], 0)                       # already recorded → skipped
        self.assertEqual(ARReceipt.objects.count(), 3)

    def test_dashboard(self):
        a = self.acct()
        ARReceipt.objects.create(account_id=a['id'], paid_on=date(2025, 8, 1), amount=D('100000'), mode='bank')
        d = self.api.get('/api/ar/dashboard/?as_of=2026-09-21').json()
        self.assertEqual(d['accounts'], 1)
        self.assertEqual(d['totals']['received'], 100000)
        self.assertEqual(d['top_over_180'][0]['client'], 'Jigar Makwana')


class ARScheduleAndReportsTests(TestCase):
    """Bookings with no installments (every Pratishtha flat), the issue flags,
    the dashboard and the printable statement."""

    def setUp(self):
        self.co = Company.objects.create(code='VIS', name='Vistara Realty')
        self.project = Project.objects.create(company=self.co, name='Pratishtha')
        self.user = User.objects.create_user('ar2@test.local', company=self.co, user_code='AR2', password='x', name='AR User',
                                             role='Employee', modules=['AR'])
        # A flat: 28,00,000 deal, no installments; 1,50,000 extra of which stamp 40,000 + reg 10,000.
        self.flat = make_booking(self.co, self.project, plot='1003', client_name='Asha Patel', installments=[],
                                 total_extra=D('150000'), stamp_duty=D('40000'), reg_fees=D('10000'),
                                 final_amount=D('2800000'))
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.id = self.api.get('/api/ar/accounts/').json()['results'][0]['id']

    def ledger(self, as_of='2026-09-21'):
        return self.api.get(f'/api/ar/accounts/{self.id}/?as_of={as_of}').json()

    def test_no_schedule_still_totals_the_whole_deal(self):
        d = self.ledger()
        self.assertTrue(d['no_schedule'])
        self.assertEqual(d['collectable'], 2750000)                    # deal − stamp − reg
        self.assertEqual(d['plan_mismatch'], 0)
        self.assertEqual([p['kind'] for p in d['plan']], ['legal', 'balance'])
        self.assertEqual(d['plan'][-1]['amount'], 2650000)             # 27,50,000 − legal 1,00,000
        self.assertEqual(d['net_interest'], 0)
        self.assertNotIn('schedule_editable', d)             # no AR-side schedule: Sales owns it

    def test_the_schedule_cannot_be_set_from_ar(self):
        self.api.patch(f'/api/ar/accounts/{self.id}/', {'schedule': [{'date': '2026-01-01', 'amount': 2650000}]}, format='json')
        self.assertTrue(self.ledger()['no_schedule'])       # ignored — only Sales can add installments

    def test_dashboard(self):
        self.api.post(f'/api/ar/accounts/{self.id}/receipts/', {'paid_on': '2026-01-05', 'amount': '500000', 'mode': 'bank'}, format='json')
        self.flat.installments = [{'no': 1, 'date': '2026-01-01', 'amt': 2650000}]   # Sales adds the schedule
        self.flat.save()
        d = self.api.get('/api/ar/dashboard/?as_of=2026-09-21').json()
        self.assertEqual(d['accounts'], 1)
        self.assertEqual(d['totals']['received'], 500000)
        self.assertEqual(d['totals']['overdue'], 2150000)
        self.assertEqual(d['overdue_accounts'], 1)
        self.assertEqual(d['top_overdue'][0]['id'], self.id)
        self.assertEqual(d['top_over_180'][0]['amount'], 2150000)
        self.assertEqual(d['projects'], [{'id': self.project.id, 'name': 'Pratishtha'}])
        self.assertEqual(d['issues'], {'no_schedule': 0, 'plan_mismatch': 0})

    def test_statement_is_print_ready_html(self):
        self.api.post(f'/api/ar/accounts/{self.id}/receipts/', {'paid_on': '2026-01-05', 'amount': '1234567', 'mode': 'cheque',
                                                                'remarks': 'HDFC <b>'}, format='json')
        r = self.api.get(f'/api/ar/accounts/{self.id}/statement/?as_of=2026-09-21')
        self.assertEqual(r.status_code, 200)
        self.assertIn('text/html', r['Content-Type'])
        html = r.content.decode()
        self.assertIn('Statement of Account', html)
        self.assertIn('Asha Patel', html)
        self.assertIn('₹ 12,34,567', html)            # Indian grouping
        self.assertNotIn('HDFC', html)                # internal remarks stay off the client's copy
        self.assertIn('Vistara Realty', html)


class AREncryptionAtRestTests(TestCase):
    """Every payment detail must be ciphertext in the database, never plaintext."""

    def test_receipt_and_account_columns_are_ciphertext(self):
        import os
        from unittest import mock
        from cryptography.fernet import Fernet
        from django.db import connection
        with mock.patch.dict(os.environ, {'FIELD_ENCRYPTION_KEY': Fernet.generate_key().decode()}):
            co = Company.objects.create(code='VIS', name='Vistara')
            project = Project.objects.create(company=co, name='Kalrav 2')
            user = User.objects.create_user('ar9@test.local', company=co, user_code='AR9', password='x', name='AR', role='Employee', modules=['AR'])
            make_booking(co, project)
            api = APIClient()
            api.force_authenticate(user)
            aid = api.get('/api/ar/accounts/').json()['results'][0]['id']
            api.post(f'/api/ar/accounts/{aid}/receipts/', {'paid_on': '2025-08-01', 'amount': '100000', 'mode': 'nbfc',
                                                          'remarks': 'HDFC 1234'}, format='json')
            api.patch(f'/api/ar/accounts/{aid}/', {'legal_due_date': '2026-06-30'}, format='json')
            with connection.cursor() as c:
                c.execute('SELECT paid_on, amount, mode, remarks FROM receivables_arreceipt')
                raw = c.fetchone()
                c.execute('SELECT legal_due_date FROM receivables_araccount')
                raw_legal = c.fetchone()[0]
            for v in (*raw, raw_legal):
                self.assertTrue(str(v).startswith('gAAAAA'), f'not encrypted: {v!r}')
            self.assertNotIn('2025-08-01', ' '.join(map(str, raw)))
            # …and the API still reads it all back, figures intact.
            d = api.get(f'/api/ar/accounts/{aid}/').json()
            self.assertEqual((d['receipts'][0]['paid_on'], d['receipts'][0]['mode'], d['received']), ('2025-08-01', 'nbfc', 100000))
            self.assertEqual(d['legal_due_date'], '2026-06-30')


class ARImportTemplateTests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(code='VIS', name='Vistara')
        self.project = Project.objects.create(company=self.co, name='Kalrav 2')
        self.user = User.objects.create_user('ar7@test.local', company=self.co, user_code='AR7', password='x', name='AR', role='Employee', modules=['AR'])
        make_booking(self.co, self.project, plot='10')
        make_booking(self.co, self.project, plot='2', client_name='Second')
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.ids = {r['plots']: r['id'] for r in self.api.get('/api/ar/accounts/').json()['results']}

    def test_template_lists_the_projects_plots(self):
        r = self.api.get(f'/api/ar/import/template/?project_id={self.project.id}')
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Payment Tracker']
        self.assertEqual([c.value for c in ws[1]], ['Plot No', 'Client', 'Paid Date', 'Paid', 'Mode', 'Remarks'])
        self.assertEqual([ws.cell(i, 1).value for i in (2, 3)], ['2', '10'])     # plot order, not text order
        # The filled-in template goes straight back through the importer.
        ws.cell(3, 3).value = date(2025, 8, 1)
        ws.cell(3, 4).value = 100000
        ws.cell(3, 5).value = 'NBFC'
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        up = SimpleUploadedFile('t.xlsx', buf.read())
        d = self.api.post('/api/ar/import/', {'file': up, 'project_id': self.project.id}, format='multipart').json()
        self.assertEqual((d['ready'], d['skipped']), (1, 0))


class ARMatchesBookingsTests(TestCase):
    """AR must carry the same deal value the Bookings page shows (less stamp and
    registration), and say why a deal appears in AR but not on that page."""

    def setUp(self):
        self.co = Company.objects.create(code='VIS', name='Vistara')
        self.project = Project.objects.create(company=self.co, name='VIP Alindra')
        self.user = User.objects.create_user('ar5@test.local', company=self.co, user_code='AR5', password='x', name='AR', role='Employee', modules=['AR'])
        self.api = APIClient()
        self.api.force_authenticate(self.user)

    def rows(self):
        return {r['booking_id']: r for r in self.api.get('/api/ar/accounts/').json()['results']}

    def test_token_only_eoi_schedule_still_carries_the_whole_deal(self):
        # EOI-19: deal 2,09,31,220; the schedule holds only the 1,00,000 token.
        b = make_booking(self.co, self.project, plot='EOI-19', final_amount=D('20931220'), total_extra=D('294220'),
                         stamp_duty=D('0'), reg_fees=D('1500'),
                         installments=[{'no': 1, 'amt': 100000, 'date': '2026-07-13', 'isNsd': True},
                                       {'no': 'Extra', 'amt': 294220, 'date': '', 'isExtra': True}])
        r = self.rows()[b.id]
        self.assertEqual(r['collectable'], 20931220 - 1500)            # deal − stamp − reg, like the Bookings page
        self.assertLess(r['plan_mismatch'], 0)                          # still flagged: the LOI schedule is short
        ledger = self.api.get(f"/api/ar/accounts/{r['id']}/").json()
        bal = [p for p in ledger['plan'] if p['kind'] == 'balance']
        self.assertEqual(len(bal), 1)
        self.assertIsNone(bal[0]['due'])

    def test_rounding_does_not_add_a_balance_line(self):
        b = make_booking(self.co, self.project, plot='5', final_amount=D('16716926'))   # schedule ₹3 short
        ledger = self.api.get(f"/api/ar/accounts/{self.rows()[b.id]['id']}/").json()
        self.assertNotIn('balance', [p['kind'] for p in ledger['plan']])

    def test_only_deals_approved_by_sales_and_accounts_are_shown(self):
        b = make_booking(self.co, self.project, plot='EOI-56')
        self.assertIn(b.id, self.rows())
        # A revision awaiting approval: the deal leaves AR, as it leaves the Bookings page…
        r1 = make_booking(self.co, self.project, plot='EOI-56', revision_of=b, revision_no=1, status='pending',
                          approval_status='REVISION R1 PENDING')
        self.assertNotIn(b.id, self.rows())
        self.assertNotIn(r1.id, self.rows())
        self.assertEqual(self.api.get('/api/ar/dashboard/').json()['accounts'], 0)
        # …and comes back on the approved revision once Sales and Accounts approve it.
        r1.status, r1.accounts_status, r1.approval_status = 'sold', 'approved', 'REVISION R1 APPROVED'
        r1.save()
        self.assertIn(r1.id, self.rows())

    def test_blank_accounts_status_is_not_approved(self):
        b = make_booking(self.co, self.project, plot='9', accounts_status='')
        self.assertNotIn(b.id, self.rows())


class AROsSummaryTests(TestCase):
    """Plot 25 (Ankitbhai Lathigara), Kalrav 2 workbook, ledger date 21/09/26: the
    O/s Summary puts the overdue into the current month and totals to the O/s."""

    def test_overdue_sits_in_the_current_month(self):
        co = Company.objects.create(code='VIS', name='Vistara')
        project = Project.objects.create(company=co, name='Kalrav 2')
        user = User.objects.create_user('ar6@test.local', company=co, user_code='AR6', password='x', name='AR', role='Employee', modules=['AR'])
        inst = [{'no': i, 'date': d, 'amt': 2158533, 'isNsd': True}
                for i, d in enumerate(['2026-07-20', '2026-08-20', '2026-09-20', '2026-10-20', '2026-11-20'], start=1)]
        make_booking(co, project, plot='25', installments=inst, total_extra=D('0'), stamp_duty=D('0'), reg_fees=D('0'),
                     final_amount=D(2158533 * 5))
        api = APIClient(); api.force_authenticate(user)
        aid = api.get('/api/ar/accounts/').json()['results'][0]['id']
        for d, a, m in (('2026-07-19', 100000, 'nbfc'), ('2026-07-29', 1000000, 'nbfc'), ('2026-08-17', 100000, 'bank'), ('2026-08-18', 1500000, 'nbfc')):
            api.post(f'/api/ar/accounts/{aid}/receipts/', {'paid_on': d, 'amount': a, 'mode': m}, format='json')
        led = api.get(f'/api/ar/accounts/{aid}/?as_of=2026-09-21').json()
        rows = {x['label']: x['amount'] for x in led['os_summary']}
        self.assertEqual(rows['Sep-26'], 3775599)                 # the workbook's Sep-26 row
        self.assertEqual(rows['Oct-26'], 2158533)
        self.assertEqual(rows['Nov-26'], 2158533)
        self.assertEqual(sum(rows.values()), led['outstanding'])
        self.assertEqual(led['net_interest'], 54501)              # the workbook's Interest Due (f); its row total
                                                                  # shows 54,502 only because each row is rounded first


class ARBookingAndLoiTests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(code='VIS', name='Vistara')
        self.project = Project.objects.create(company=self.co, name='Anahata Florenza')
        self.user = User.objects.create_user('ar8@test.local', company=self.co, user_code='AR8', password='x', name='AR', role='Employee', modules=['AR'])
        self.b = make_booking(self.co, self.project, plot='Ananda1', client_name='Yashpalsinh Thakor')
        self.api = APIClient(); self.api.force_authenticate(self.user)
        self.aid = self.api.get('/api/ar/accounts/').json()['results'][0]['id']

    def test_booking_details_for_an_ar_user(self):
        r = self.api.get(f'/api/ar/accounts/{self.aid}/booking/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['client_name'], 'Yashpalsinh Thakor')
        self.assertEqual(len(r.json()['installments']), 5)

    def test_loi_link_says_why_when_there_is_none(self):
        r = self.api.get(f'/api/ar/accounts/{self.aid}/loi-url/')
        self.assertEqual(r.status_code, 404)
        self.assertIn('No LOI', r.json()['detail'])

    def test_no_ar_access_no_details(self):
        other = User.objects.create_user('no8@test.local', company=self.co, user_code='NO8', password='x', name='N', role='Employee', modules=['Sales'])
        c = APIClient(); c.force_authenticate(other)
        self.assertEqual(c.get(f'/api/ar/accounts/{self.aid}/booking/').status_code, 403)
        self.assertEqual(c.get(f'/api/ar/accounts/{self.aid}/loi-url/').status_code, 403)
