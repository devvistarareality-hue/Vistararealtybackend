"""Ask Nexora: the query tool reads only what the asker may see, handles encrypted
fields, and the question loop runs tools then answers (Claude replaced by a fake)."""
from types import SimpleNamespace
from unittest import mock

from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import User
from activity.models import ActivityLog
from companies.models import Company
from sales import assistant
from sales.models import Closure, Lead, LeadSource, Project, SiteVisit


class QueryToolTests(TestCase):
    def setUp(self):
        cache.clear()
        self.co = Company.objects.create(code='AIQ', name='Aiq Co')
        self.p1 = Project.objects.create(company=self.co, name='Kalrav')
        self.p2 = Project.objects.create(company=self.co, name='Tundav')
        meta = LeadSource.objects.create(company=self.co, name='Meta')
        mk = lambda code, role, boss=None: User.objects.create(
            name=code, email=f'{code}@x.com', phone='9' + code + '00000', user_code=code, role=role,
            designation='STM' if role == 'Employee' else role, company=self.co, modules=['Sales'],
            reporting_manager=boss)
        self.admin = mk('QA', 'Admin')
        self.boss = mk('QB', 'Manager')
        self.stm = mk('QS', 'Employee', self.boss)
        self.other = mk('QO', 'Employee', self.boss)
        now = timezone.now()
        for i, (owner, proj) in enumerate([(self.stm, self.p1), (self.stm, self.p2), (self.other, self.p1)]):
            lead = Lead.objects.create(company=self.co, project=proj, source=meta, name=f'Shah {i}',
                                       phone=f'98000000{i:02d}', stm=owner, meta_campaign_name='Diwali')
            SiteVisit.objects.create(lead=lead, project=proj, stm=owner, status='completed',
                                     outcome='hot' if i == 0 else 'cold', scheduled_at=now, visited_at=now)
            Closure.objects.create(company=self.co, lead=lead, project=proj, stm=owner, client_name=f'Shah {i}',
                                   closure_date=timezone.localdate(), total_amount=1000000 * (i + 1))

    def q(self, user, **args):
        return assistant.run_query(user, None, args)

    def test_counts_are_limited_to_what_the_asker_sees(self):
        self.assertEqual(self.q(self.stm, entity='leads', mode='count')['total'], 2)
        self.assertEqual(self.q(self.admin, entity='leads', mode='count')['total'], 3)

    def test_group_by_and_date_keywords(self):
        out = self.q(self.admin, entity='site_visits', mode='count', group_by=['project'],
                     filters=[{'field': 'visited', 'op': 'eq', 'value': 'today'}])
        self.assertEqual({g['project']: g['count'] for g in out['groups']}, {'Kalrav': 2, 'Tundav': 1})

    def test_encrypted_filter_and_money_sum(self):
        out = self.q(self.admin, entity='leads', mode='count',
                     filters=[{'field': 'name', 'op': 'contains', 'value': 'shah 1'}])
        self.assertEqual(out['total'], 1)
        out = self.q(self.admin, entity='closures', mode='sum', sum_field='total_amount', group_by=['project'])
        self.assertEqual(out['grand_total'], 6000000.0)
        self.assertEqual({g['project']: g['total'] for g in out['groups']}, {'Kalrav': 4000000.0, 'Tundav': 2000000.0})

    def test_list_mode(self):
        out = self.q(self.stm, entity='site_visits', mode='list', fields=['client', 'project', 'outcome'])
        self.assertEqual(out['shown'], 2)
        self.assertTrue(all(set(r) == {'client', 'project', 'outcome'} for r in out['rows']))

    def test_unknown_field_is_an_error(self):
        with self.assertRaises(ValueError):
            self.q(self.admin, entity='leads', mode='count', group_by=['name'])


def _msg(content, stop):
    usage = SimpleNamespace(input_tokens=1000, output_tokens=200, cache_read_input_tokens=0,
                            cache_creation_input_tokens=0)
    return SimpleNamespace(stop_reason=stop, usage=usage, to_dict=lambda: {'content': content})


@override_settings(AI_INLINE=True)
class AskViewTests(TestCase):
    def setUp(self):
        cache.clear()
        self.co = Company.objects.create(code='AIV', name='Aiv Co')
        self.admin = User.objects.create(name='Admin', email='a@aiv.com', phone='9100000099', user_code='AV-A',
                                         role='Admin', company=self.co)
        self.stm = User.objects.create(name='S', email='s@aiv.com', phone='9100000098', user_code='AV-S',
                                       role='Employee', designation='STM', company=self.co, modules=['Sales'],
                                       reporting_manager=self.admin)

    def _ask(self, user, question='How many leads?'):
        c = APIClient(); c.force_authenticate(user)
        res = c.post('/api/sales/ai/ask/', {'question': question}, format='json')
        return c, res

    def test_not_ticked_is_refused(self):
        _, res = self._ask(self.stm)
        self.assertEqual(res.status_code, 403)

    @mock.patch.dict('os.environ', {'ANTHROPIC_API_KEY': 'test'})
    def test_tool_then_answer_and_logged(self):
        fake = mock.MagicMock()
        fake.messages.create.side_effect = [
            _msg([{'type': 'tool_use', 'id': 't1', 'name': 'query_records',
                   'input': {'entity': 'leads', 'mode': 'count'}}], 'tool_use'),
            _msg([{'type': 'text', 'text': 'You have 0 leads.'}], 'end_turn'),
        ]
        with mock.patch.object(assistant, '_client', return_value=fake):
            c, res = self._ask(self.admin)
        self.assertEqual(res.status_code, 202, res.content)
        st = c.get(f"/api/sales/ai/ask/{res.json()['job']}/").json()
        self.assertEqual(st['status'], 'done', st)
        self.assertEqual(st['answer'], 'You have 0 leads.')
        # the tool result went back to the model
        sent = fake.messages.create.call_args_list[1].kwargs['messages']
        result = next(m for m in sent if m['role'] == 'user' and isinstance(m['content'], list))
        self.assertEqual(result['content'][0]['type'], 'tool_result')
        self.assertIn('"total": 0', result['content'][0]['content'])
        import json
        log = json.loads(ActivityLog.objects.filter(target_type='ai question').latest('id').details)
        self.assertEqual(log['queries'], 1)
        self.assertGreater(log['approx_cost_inr'], 0)
        # someone else cannot read the answer
        other = APIClient(); other.force_authenticate(self.stm)
        self.assertEqual(other.get(f"/api/sales/ai/ask/{res.json()['job']}/").status_code, 404)

    def test_missing_key_says_so(self):
        with mock.patch.dict('os.environ', {}, clear=False):
            import os
            os.environ.pop('ANTHROPIC_API_KEY', None)
            _, res = self._ask(self.admin)
        self.assertEqual(res.status_code, 503)


class OtherModuleTests(TestCase):
    """Ask Nexora reads every module the asker holds, and only those."""
    def setUp(self):
        cache.clear()
        from tasks.models import Task, TaskList
        from attendance.models import LeaveApplication
        self.co = Company.objects.create(code='AIM', name='Aim Co')
        self.admin = User.objects.create(name='Boss', email='b@aim.com', phone='9200000001', user_code='AM-A',
                                         role='Admin', company=self.co)
        self.stm = User.objects.create(name='Stm', email='s@aim.com', phone='9200000002', user_code='AM-S',
                                       role='Employee', designation='STM', company=self.co, modules=['Sales'],
                                       reporting_manager=self.admin)
        tl = TaskList.objects.create(company=self.co, name='Site work')
        for i, st in enumerate(['todo', 'done', 'todo']):
            t = Task.objects.create(company=self.co, task_list=tl, title=f'T{i}', code=f'TSK-00{i}', status=st,
                                    due_date=timezone.localdate())
            t.assignees.set([self.stm])
        LeaveApplication.objects.create(user=self.stm, leave_type='sick_leave', from_date=timezone.localdate())

    def test_modules_follow_user_management(self):
        self.assertIn('tasks', assistant.available(self.admin))
        self.assertIn('ar_accounts', assistant.available(self.admin))
        got = assistant.available(self.stm)
        self.assertIn('leads', got)
        self.assertIn('my_leaves', got)          # everyone has their own HR records
        self.assertNotIn('tasks', got)
        self.assertNotIn('ar_accounts', got)
        with self.assertRaises(ValueError):
            assistant.run_query(self.stm, None, {'entity': 'tasks', 'mode': 'count'})
        tool = assistant.tool_for(got)
        self.assertNotIn('ar_accounts', tool['description'])

    def test_tasks_and_leaves(self):
        out = assistant.run_query(self.admin, None, {'entity': 'tasks', 'mode': 'count', 'group_by': ['status']})
        self.assertEqual({g['status']: g['count'] for g in out['groups']}, {'todo': 2, 'done': 1})
        self.assertEqual(assistant.run_query(self.admin, None, {'entity': 'team_leaves', 'mode': 'count'})['total'], 1)
        self.assertEqual(assistant.run_query(self.stm, None, {'entity': 'my_leaves', 'mode': 'count'})['total'], 1)
        self.assertEqual(assistant.run_query(self.stm, None, {'entity': 'team_leaves', 'mode': 'count'})['total'], 0)

    def test_computed_rows(self):
        rows = [{'project': 'Kalrav', 'overdue': '1000.00', 'booking_date': '2026-09-01'},
                {'project': 'Kalrav', 'overdue': '500.00', 'booking_date': '2026-10-01'},
                {'project': 'Tundav', 'overdue': '0', 'booking_date': None}]
        fields = assistant._entities()['ar_accounts']['fields']
        out = assistant._run_rows(rows, fields, {'mode': 'sum', 'sum_field': 'overdue', 'group_by': ['project'],
                                                 'filters': [{'field': 'overdue', 'op': 'gte', 'value': 1}]},
                                  'booking_date')
        self.assertEqual(out['grand_total'], 1500.0)
        self.assertEqual(out['groups'], [{'project': 'Kalrav', 'total': 1500.0}])
        out = assistant._run_rows(rows, fields, {'mode': 'count', 'group_by': ['booking_date:month']}, 'booking_date')
        self.assertEqual(out['total'], 3)
        # The AR register hands its computed rows over (an admin has AR).
        self.assertEqual(assistant.run_query(self.admin, None, {'entity': 'ar_accounts', 'mode': 'count'})['total'], 0)


@override_settings(AI_INLINE=True)
class AiFileTests(TestCase):
    """Ask Nexora makes Excel / PDF files: full query data, the asker's scope, lists
    only with the Download Excel tick, and only the asker can fetch the file."""
    def setUp(self):
        cache.clear()
        self.co = Company.objects.create(code='AIF', name='Aif Co')
        self.p = Project.objects.create(company=self.co, name='Kalrav')
        self.admin = User.objects.create(name='Boss', email='b@aif.com', phone='9500000001', user_code='AF-A',
                                         role='Admin', company=self.co)
        self.stm = User.objects.create(name='Stm', email='s@aif.com', phone='9500000002', user_code='AF-S',
                                       role='Employee', designation='STM', company=self.co, modules=['Sales'],
                                       can_use_ai=True, reporting_manager=self.admin)
        for i in range(3):
            Lead.objects.create(company=self.co, project=self.p, name=f'Shah {i}', phone=f'95000000{i:02d}',
                                stm=self.stm)

    def test_excel_and_pdf_from_queries(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from sales import ai_files
        spec = {'format': 'xlsx', 'title': 'Leads ₹ report', 'summary': 'Three leads.',
                'tables': [{'heading': 'By project', 'query': {'entity': 'leads', 'mode': 'count', 'group_by': ['project']}},
                           {'heading': 'All leads', 'query': {'entity': 'leads', 'mode': 'list', 'fields': ['name', 'project']}}]}
        made = ai_files.make(self.admin, None, spec)
        path, name, mime = ai_files.read(made['id'], self.admin)
        self.assertTrue(name.endswith('.xlsx'))
        with open(path, 'rb') as f:
            wb = load_workbook(BytesIO(f.read()))
        self.assertEqual(wb.sheetnames, ['Summary', 'By project', 'All leads'])
        self.assertEqual(wb['All leads'].max_row, 4)                 # header + 3 leads
        self.assertIsNone(ai_files.read(made['id'], self.stm))       # not theirs
        pdf = ai_files.make(self.admin, None, {**spec, 'format': 'pdf'})
        with open(ai_files.read(pdf['id'], self.admin)[0], 'rb') as f:
            self.assertTrue(f.read(5).startswith(b'%PDF'))

    def test_lists_need_the_download_tick(self):
        from sales import ai_files
        spec = {'format': 'xlsx', 'title': 'x',
                'tables': [{'heading': 'L', 'query': {'entity': 'leads', 'mode': 'list'}}]}
        with self.assertRaises(PermissionError):
            ai_files.make(self.stm, None, spec)
        summary = {'format': 'xlsx', 'title': 'x',
                   'tables': [{'heading': 'C', 'query': {'entity': 'leads', 'mode': 'count'}}]}
        self.assertTrue(ai_files.make(self.stm, None, summary)['id'])

    @mock.patch.dict('os.environ', {'ANTHROPIC_API_KEY': 'test'})
    def test_file_reaches_the_panel_and_downloads(self):
        fake = mock.MagicMock()
        fake.messages.create.side_effect = [
            _msg([{'type': 'tool_use', 'id': 'f1', 'name': 'make_file', 'input': {
                'format': 'xlsx', 'title': 'Lead count',
                'tables': [{'heading': 'Leads', 'query': {'entity': 'leads', 'mode': 'count'}}]}}], 'tool_use'),
            _msg([{'type': 'text', 'text': 'Your Excel is ready.'}], 'end_turn'),
        ]
        c = APIClient(); c.force_authenticate(self.admin)
        with mock.patch.object(assistant, '_client', return_value=fake):
            res = c.post('/api/sales/ai/ask/', {'question': 'make an excel of my leads'}, format='json')
        st = c.get(f"/api/sales/ai/ask/{res.json()['job']}/").json()
        self.assertEqual(st['status'], 'done', st)
        self.assertEqual(len(st['files']), 1)
        dl = c.get(f"/api/sales/ai/files/{st['files'][0]['id']}/")
        self.assertEqual(dl.status_code, 200)
        other = APIClient(); other.force_authenticate(self.stm)
        self.assertEqual(other.get(f"/api/sales/ai/files/{st['files'][0]['id']}/").status_code, 404)


@override_settings(AI_INLINE=True)
class ChatHistoryTests(TestCase):
    """Chats are saved (encrypted), private to their owner, continue with their own
    history, and every question is logged under the Ask Nexora module."""
    def setUp(self):
        cache.clear()
        self.co = Company.objects.create(code='AIC', name='Aic Co')
        self.admin = User.objects.create(name='Boss', email='b@aic.com', phone='9600000001', user_code='AC-A',
                                         role='Admin', company=self.co)
        self.other = User.objects.create(name='Other', email='o@aic.com', phone='9600000002', user_code='AC-O',
                                         role='Admin', company=self.co)

    @mock.patch.dict('os.environ', {'ANTHROPIC_API_KEY': 'test'})
    def test_chat_is_saved_and_continued(self):
        from sales.models import AiChat
        fake = mock.MagicMock()
        fake.messages.create.side_effect = [
            _msg([{'type': 'text', 'text': 'First answer.'}], 'end_turn'),
            _msg([{'type': 'text', 'text': 'Second answer.'}], 'end_turn'),
        ]
        c = APIClient(); c.force_authenticate(self.admin)
        with mock.patch.object(assistant, '_client', return_value=fake):
            r1 = c.post('/api/sales/ai/ask/', {'question': 'How many leads?'}, format='json').json()
            r2 = c.post('/api/sales/ai/ask/', {'question': 'And last month?', 'chat_id': r1['chat']}, format='json').json()
        self.assertEqual(r1['chat'], r2['chat'])
        # the second call carried the first turn as history
        sent = fake.messages.create.call_args_list[1].kwargs['messages']
        self.assertEqual(sent[0]['content'], 'How many leads?')
        lst = c.get('/api/sales/ai/chats/').json()['results']
        self.assertEqual([(x['title'], x['turns']) for x in lst], [('How many leads?', 2)])
        detail = c.get(f"/api/sales/ai/chats/{r1['chat']}/").json()
        self.assertEqual([t['a'] for t in detail['turns']], ['First answer.', 'Second answer.'])
        # encrypted at rest
        from django.db import connection
        with connection.cursor() as cur:
            cur.execute('SELECT turns FROM sales_aichat WHERE id = %s', [r1['chat']])
            self.assertNotIn('First answer', cur.fetchone()[0])
        # private
        o = APIClient(); o.force_authenticate(self.other)
        self.assertEqual(o.get(f"/api/sales/ai/chats/{r1['chat']}/").status_code, 404)
        self.assertEqual(o.get('/api/sales/ai/chats/').json()['results'], [])
        # logged under Ask Nexora
        self.assertEqual(set(ActivityLog.objects.filter(target_type='ai question').values_list('module', flat=True)),
                         {'Ask Nexora'})
        self.assertEqual(c.delete(f"/api/sales/ai/chats/{r1['chat']}/").status_code, 204)
        self.assertFalse(AiChat.objects.exists())


class ScopeLeakTests(TestCase):
    """A CP-only person reads the partner book only; a platform admin's answers stay
    in one company."""
    def test_cp_only_user_is_confined_to_the_cp_book(self):
        from sales.models import ChannelPartner
        co = Company.objects.create(code='SLK', name='Slk')
        p = Project.objects.create(company=co, name='Kalrav')
        cp = ChannelPartner.objects.create(company=co, name='Broker', contact_no='9700000001')
        mgr = User.objects.create(name='CPM', email='m@slk.com', phone='9700000010', user_code='SL-M', role='Manager',
                                  designation='CP Cluster Head', company=co, modules=['Channel Partner'],
                                  manager_modules=['Channel Partner'])
        Lead.objects.create(company=co, project=p, name='Sales lead', phone='9700000020')
        Lead.objects.create(company=co, project=p, name='CP lead', phone='9700000021', channel_partner=cp)
        self.assertEqual(assistant._book_for(mgr, 'all'), 'cp')
        names = [r['name'] for r in assistant.run_query(mgr, None, {'entity': 'leads', 'mode': 'list', 'fields': ['name'],
                                                                    'book': 'all'})['rows']]
        self.assertNotIn('Sales lead', names)
        from sales.models import Booking
        stm = User.objects.create(name='STM', email='s@slk.com', phone='9700000011', user_code='SL-S', role='Employee',
                                  designation='STM', company=co, modules=['Sales'], reporting_manager=mgr)
        # Someone else's Sales deal (the CP screen also shows the manager's own deals, any source).
        Booking.objects.create(company=co, project=p, client_name='Sales deal', status='sold', source='Walk-In', stm=stm)
        Booking.objects.create(company=co, project=p, client_name='CP deal', status='sold', source='Channel Partner', stm=mgr)
        deals = [r['client'] for r in assistant.run_query(mgr, None, {'entity': 'bookings', 'mode': 'list',
                                                                      'fields': ['client'], 'book': 'all'})['rows']]
        self.assertNotIn('Sales deal', deals)
