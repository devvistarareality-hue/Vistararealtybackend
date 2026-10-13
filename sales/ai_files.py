"""Files Ask Nexora makes on request — "make an Excel of…", "give me a PDF report of…".

The model calls the make_file tool (sales/assistant.py) with a title, a short
summary, a few notes and up to eight tables. A table is either rows the model
writes itself (a small comparison it worked out) or — far more often — a
query_records query, run here in full so the file carries the real data rather
than what the model retyped (and up to FILE_ROWS rows, not the chat's 100).

Who may have what: the same scoping as the screens (every query goes through the
asker's own list views), and a file that lists records — names, phones, units —
needs the Download Excel tick in User Management, exactly like the Leads and Site
Visits downloads. Summary tables (counts, totals) are open to anyone with Ask Nexora.

Files are written to a folder every worker of the container shares, readable only
by the person who asked, and kept for a day.
"""
import json
import os
import time
import uuid
from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone

FILE_DIR = os.path.join('/tmp', 'nexora-ai', 'files')
MAX_AGE = 24 * 60 * 60
FILE_ROWS = 5000          # rows a list table may carry into an Excel
PDF_ROWS = 400            # … and into a PDF (pages, not a database)
MAX_TABLES = 8
MIMES = {'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'pdf': 'application/pdf'}

TOOL = {
    'name': 'make_file',
    'description': (
        "Make a downloadable Excel (xlsx) or PDF report when the person asks for a file, sheet, "
        "Excel, PDF, report or download. Query first to understand the data, then call this once. "
        "Prefer tables with `query` (the same arguments as query_records) so the file holds the full, "
        "exact data — for a list use mode 'list' with the fields to include; for a breakdown use "
        "count/sum with group_by. Use `rows` only for a small table you computed yourself. After it "
        "returns, reply in one or two lines that the file is ready (a download button appears) and "
        "what is in it."
    ),
    'input_schema': {
        'type': 'object',
        'properties': {
            'format': {'type': 'string', 'enum': ['xlsx', 'pdf']},
            'title': {'type': 'string', 'description': 'Report title, e.g. "AR overdue by project — Oct 2026".'},
            'summary': {'type': 'string', 'description': 'Two to five sentences: the headline numbers.'},
            'notes': {'type': 'array', 'items': {'type': 'string'}, 'maxItems': 8,
                      'description': 'Key points or next steps, one line each.'},
            'tables': {'type': 'array', 'maxItems': MAX_TABLES, 'items': {
                'type': 'object',
                'properties': {
                    'heading': {'type': 'string'},
                    'query': {'type': 'object', 'description': 'query_records arguments for this table.'},
                    'columns': {'type': 'array', 'items': {'type': 'string'}},
                    'rows': {'type': 'array', 'items': {'type': 'array'}},
                },
                'required': ['heading'],
            }},
        },
        'required': ['format', 'title'],
    },
}


def can_list(user):
    """A file with record lists needs the Download Excel tick (or a real admin)."""
    from .views import _is_hard_admin
    return bool(_is_hard_admin(user) or getattr(user, 'can_export_leads', False)
                or getattr(user, 'role', '') == 'Admin')


def _label(name):
    return str(name).replace('_', ' ').replace(':', ' by ').strip().capitalize()


def _table_from_result(out, args):
    """A query_records result → (columns, rows)."""
    if 'rows' in out:
        cols = list(out['rows'][0].keys()) if out['rows'] else [f for f in (args.get('fields') or [])]
        return [_label(c) for c in cols], [[r.get(c) for c in cols] for r in out['rows']]
    groups = args.get('group_by') or []
    if out.get('groups'):
        val = 'total' if args.get('mode') == 'sum' else 'count'
        head = [_label(g) for g in groups] + [_label(args.get('sum_field')) if val == 'total' else 'Count']
        rows = [[g.get(k) for k in groups] + [g.get(val)] for g in out['groups']]
        total = out.get('grand_total', out.get('total'))
        if total is not None:
            rows.append(['Total'] + [''] * (len(groups) - 1) + [total])
        return head, rows
    if args.get('mode') == 'sum':
        return [_label(args.get('sum_field') or 'Total')], [[out.get('grand_total')]]
    return ['Count'], [[out.get('total')]]


def build_tables(user, company_id, spec, fmt):
    """Run every table's query as the asker; returns [(heading, columns, rows)]."""
    from .assistant import run_query
    cap = FILE_ROWS if fmt == 'xlsx' else PDF_ROWS
    tables = []
    for t in (spec.get('tables') or [])[:MAX_TABLES]:
        heading = str(t.get('heading') or 'Table')[:80]
        if isinstance(t.get('query'), dict):
            args = dict(t['query'])
            if (args.get('mode') or 'count') == 'list':
                if not can_list(user):
                    raise PermissionError('This person may not download record lists (no Download Excel tick in '
                                          'User Management). Offer a summary file with counts and totals instead.')
                args['limit'] = cap
            out = run_query(user, company_id, args, list_cap=cap, mask=False)   # files keep real numbers
            cols, rows = _table_from_result(out, args)
        else:
            cols = [str(c) for c in (t.get('columns') or [])][:30]
            rows = [list(r)[:30] for r in (t.get('rows') or [])][:cap]
        tables.append((heading, cols, rows))
    return tables


def _cell(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v if isinstance(v, (int, float)) else ('' if v is None else str(v))


def _xlsx(spec, tables, who):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = spec.get('title') or 'Report'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = who
    ws['A2'].font = Font(italic=True, color='666666')
    r = 4
    if spec.get('summary'):
        ws.cell(r, 1, str(spec['summary'])).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 15 * max(2, len(str(spec['summary'])) // 90 + 1)
        r += 2
    for n in spec.get('notes') or []:
        ws.cell(r, 1, '• ' + str(n))
        r += 1
    ws.column_dimensions['A'].width = 110
    head_fill = PatternFill('solid', fgColor='1F3A5F')
    used = {'Summary'}
    for heading, cols, rows in tables:
        name = ''.join(c for c in heading if c not in '[]:*?/\\')[:28] or 'Table'
        base, i = name, 2
        while name in used:
            name, i = f'{base[:25]} {i}', i + 1
        used.add(name)
        sh = wb.create_sheet(name)
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head_fill
        for row in rows:
            sh.append([_cell(v) for v in row])
        sh.freeze_panes = 'A2'
        for idx, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [len(str(rw[idx - 1])) for rw in rows[:300] if idx - 1 < len(rw)])
            sh.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 50)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_text(s):
    # The built-in PDF fonts have no rupee sign.
    return escape(str(s if s is not None else '').replace('₹', 'Rs. '))


def _pdf(spec, tables, who):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    wide = any(len(cols) > 6 for _, cols, _ in tables)
    size = landscape(A4) if wide else A4
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=spec.get('title') or 'Report')
    st = getSampleStyleSheet()
    h1 = ParagraphStyle('h1', parent=st['Title'], fontSize=17, alignment=0, spaceAfter=4)
    meta = ParagraphStyle('meta', parent=st['Normal'], fontSize=8.5, textColor=colors.HexColor('#667085'))
    body = ParagraphStyle('body', parent=st['Normal'], fontSize=10, leading=14)
    h2 = ParagraphStyle('h2', parent=st['Heading3'], fontSize=11.5, spaceBefore=10, spaceAfter=5)
    cellst = ParagraphStyle('cell', parent=st['Normal'], fontSize=8, leading=10)
    headst = ParagraphStyle('head', parent=cellst, textColor=colors.white, fontName='Helvetica-Bold')
    story = [Paragraph(_pdf_text(spec.get('title') or 'Report'), h1), Paragraph(_pdf_text(who), meta), Spacer(1, 8)]
    if spec.get('summary'):
        story += [Paragraph(_pdf_text(spec['summary']), body), Spacer(1, 6)]
    for n in spec.get('notes') or []:
        story.append(Paragraph('• ' + _pdf_text(n), body))
    width = size[0] - 28 * mm
    for heading, cols, rows in tables:
        story.append(Paragraph(_pdf_text(heading), h2))
        if not cols:
            continue
        data = [[Paragraph(_pdf_text(c), headst) for c in cols]]
        data += [[Paragraph(_pdf_text(_cell(v))[:300], cellst) for v in row] for row in rows]
        t = Table(data, repeatRows=1, colWidths=[width / len(cols)] * len(cols))
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3A5F')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6FA')]),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#D0D7E2')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t)
    story += [Spacer(1, 10), Paragraph('Made by Ask Nexora from the data this person can see. Check important '
                                       'numbers on the screens.', meta)]
    doc.build(story)
    return buf.getvalue()


def make(user, company_id, spec):
    """Build and store the file; returns {'id', 'name', 'format', 'rows'} for the model and the panel."""
    fmt = 'pdf' if spec.get('format') == 'pdf' else 'xlsx'
    tables = build_tables(user, company_id, spec, fmt)
    who = 'Prepared for %s · %s · Nexora' % (user.name or 'you', timezone.localtime().strftime('%d %b %Y, %I:%M %p'))
    data = (_pdf if fmt == 'pdf' else _xlsx)(spec, tables, who)
    title = str(spec.get('title') or 'Nexora report')
    safe = ''.join(c if c.isalnum() or c in ' -_' else ' ' for c in title).strip()[:60] or 'Nexora report'
    name = '%s %s.%s' % (safe, timezone.localdate().isoformat(), fmt)
    fid = uuid.uuid4().hex
    os.makedirs(FILE_DIR, exist_ok=True)
    with open(os.path.join(FILE_DIR, fid + '.' + fmt), 'wb') as f:
        f.write(data)
    with open(os.path.join(FILE_DIR, fid + '.json'), 'w') as f:
        json.dump({'owner': user.id, 'name': name, 'format': fmt, 'at': time.time()}, f)
    return {'id': fid, 'name': name, 'format': fmt, 'rows': sum(len(r) for _, _, r in tables)}


def read(fid, user):
    """(path, name, mime) for the asker's own file, or None."""
    fid = ''.join(c for c in str(fid) if c.isalnum())
    try:
        with open(os.path.join(FILE_DIR, fid + '.json')) as f:
            meta = json.load(f)
    except (OSError, ValueError):
        return None
    path = os.path.join(FILE_DIR, fid + '.' + meta.get('format', 'xlsx'))
    if meta.get('owner') != user.id or time.time() - meta.get('at', 0) > MAX_AGE or not os.path.exists(path):
        return None
    return path, meta['name'], MIMES.get(meta.get('format'), 'application/octet-stream')
