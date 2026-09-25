"""Excel backup of one company: a workbook with a sheet per module, and the
restore that reads it back.

Built around one workflow. An admin takes a snapshot, runs Data Reset
(SalesDataResetView) to wipe that company's trial Sales data, then uploads the
same file to put it back — with the original row ids, so everything that
pointed at a record still resolves.

Only the tables Data Reset actually clears are restorable. The rest (projects,
plots, users, AR, tasks, Club 1000 …) are exported so the snapshot is complete,
but never written back: the reset never removed them, so there is nothing to
put back and overwriting live rows is exactly what we don't want.

Encrypted columns need no special handling here. EncryptedTextField and friends
(see fields.py) decrypt on read and encrypt on write, so reading an attribute
gives plaintext for the sheet and writing plaintext back re-encrypts it.

A workbook only ever goes back into the company it came from. The request being
scoped to one company decides which file you get, not which rows get written —
so restore checks the rows themselves and refuses a file holding anyone else's.
"""
import json
import uuid
from contextlib import contextmanager
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

import openpyxl
from django.apps import apps
from django.db import models, transaction
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill

NAVY = 'FF0F1838'

# Never written to a sheet. Password hashes and session tokens would be a
# liability sitting in a downloaded file, and the email blind index is a
# derived lookup key nobody reading a backup needs.
SENSITIVE = {'password', 'session_token_app', 'session_token_web', 'email_key'}

# Marks the readable half of a foreign key pair. Restore skips these columns —
# the id column next to it is the real value.
LABEL_SUFFIX = ' (name)'


class Table:
    """One stacked table inside a module sheet."""

    def __init__(self, label, model, scope, restorable=False):
        self.label = label
        self.model = model        # 'app_label.ModelName'
        self.scope = scope        # ORM path from this model to companies.Company
        self.restorable = restorable

    @property
    def cls(self):
        return apps.get_model(self.model)


# Sheet -> the tables stacked inside it. Within a sheet the order is also the
# order restore writes them in, so a table never refers to one below it:
# closures before bookings (Booking.closure), leads before both.
SHEETS = [
    ('Sales', [
        # Restore writes these in this order, so it follows the pipeline: a lead,
        # then its visits, then the closure that can name one, then the booking
        # that names the closure.
        Table('Leads', 'sales.Lead', 'company', restorable=True),
        Table('Site Visits', 'sales.SiteVisit', 'lead__company', restorable=True),
        Table('Closures', 'sales.Closure', 'company', restorable=True),
        Table('Bookings', 'sales.Booking', 'company', restorable=True),
        Table('Follow-Ups', 'sales.FollowUp', 'lead__company', restorable=True),
        Table('Lead History', 'sales.LeadStatusHistory', 'lead__company', restorable=True),
        Table('Distribution Log', 'sales.DistributionLog', 'company', restorable=True),
        Table('Availability', 'sales.UserAvailability', 'user__company', restorable=True),
        Table('Notifications', 'accounts.Notification', 'recipient__company', restorable=True),
        # Data Reset never names lead transfers, but LeadTransfer.lead is CASCADE,
        # so wiping leads takes them with it — restorable or they are lost for good.
        Table('Lead Transfers', 'sales.LeadTransfer', 'company', restorable=True),
        # Data Reset keeps these, so they are snapshot-only.
        Table('Lead Sources', 'sales.LeadSource', 'company'),
        Table('Projects', 'sales.Project', 'company'),
        Table('Plots', 'sales.Plot', 'project__company'),
        Table('Project Assignments', 'sales.UserProjectAssignment', 'user__company'),
        Table('Sales Team', 'sales.SalesTeamMember', 'user__company'),
        Table('Distribution Settings', 'sales.DistributionSettings', 'company'),
        Table('Distribution Weights', 'sales.UserDistributionWeight', 'user__company'),
        Table('Meta Form Mappings', 'sales.MetaFormMapping', 'company'),
        Table('Meta Webhook Config', 'sales.MetaWebhookConfig', 'company'),
    ]),
    ('Channel Partner', [
        Table('Channel Partners', 'sales.ChannelPartner', 'company'),
    ]),
    ('HR', [
        Table('Users', 'accounts.User', 'company'),
        Table('Designations', 'accounts.Designation', 'company'),
        Table('Role Dashboards', 'accounts.RoleDashboard', 'company'),
        Table('Attendance', 'attendance.AttendanceRecord', 'user__company'),
        Table('Leave Applications', 'attendance.LeaveApplication', 'user__company'),
        Table('Leave Balances', 'attendance.LeaveBalance', 'user__company'),
        Table('Leave Transactions', 'attendance.LeaveTransaction', 'user__company'),
    ]),
    ('AR', [
        Table('AR Accounts', 'receivables.ARAccount', 'company'),
        Table('AR Receipts', 'receivables.ARReceipt', 'account__company'),
        Table('AR Follow-Ups', 'receivables.ARFollowUp', 'account__company'),
        Table('AR Receipt Audit', 'receivables.ARReceiptAudit', 'receipt__account__company'),
    ]),
    ('Task Allocation', [
        Table('Task Lists', 'tasks.TaskList', 'company'),
        Table('Task Tags', 'tasks.TaskTag', 'company'),
        Table('Tasks', 'tasks.Task', 'company'),
        Table('Task Checklist', 'tasks.TaskChecklistItem', 'task__company'),
        Table('Task Comments', 'tasks.TaskComment', 'task__company'),
    ]),
    ('Club 1000', [
        Table('Club Leads', 'club1000.Lead', 'company'),
        Table('Club Lead History', 'club1000.LeadStatusHistory', 'lead__company'),
        Table('Club Follow-Ups', 'club1000.FollowUp', 'lead__company'),
        Table('Schemes', 'club1000.Scheme', 'company'),
        Table('Investors', 'club1000.Investor', 'company'),
        Table('Payouts', 'club1000.Payout', 'investor__company'),
        Table('Referral Rewards', 'club1000.ReferralReward', 'investor__company'),
    ]),
    ('Activity Log', [
        Table('Activity Log', 'activity.ActivityLog', 'company'),
    ]),
]

ALL_TABLES = [t for _, tables in SHEETS for t in tables]
RESTORABLE = [t for t in ALL_TABLES if t.restorable]
ALL_LABELS = {t.label for t in ALL_TABLES}


# ── Columns ──────────────────────────────────────────────────────────────────
def _columns(model):
    """[(header, field, kind)] for one model.

    A foreign key gets two columns: the raw id, which is what restore reads,
    and a readable label beside it, which is what a person reading the sheet
    wants. Many-to-many fields are not in _meta.fields and are not exported —
    none of the restorable tables has one.
    """
    cols = []
    for f in model._meta.fields:
        if f.name in SENSITIVE:
            continue
        if f.many_to_one or f.one_to_one:
            cols.append((f.attname, f, 'id'))
            cols.append((f.name + LABEL_SUFFIX, f, 'label'))
        else:
            cols.append((f.name, f, 'value'))
    return cols


def _cell(value):
    """Coerce a Python value into something openpyxl will write."""
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        # ISO text, not an Excel date cell. Excel keeps dates as floating-point
        # days, which rounds a timestamp to the nearest millisecond and loses the
        # offset — a backup that comes back a few microseconds off is not the
        # same backup. ISO keeps both exactly, reads fine, and still sorts
        # chronologically because it is zero-padded most-significant-first.
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, date):
        return value          # a whole day survives Excel's date cell exactly
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)
    if isinstance(value, uuid.UUID):
        return str(value)
    return str(value)


# ── Export ───────────────────────────────────────────────────────────────────
# A readable label beside each foreign key is worth having, but asking each row
# for str(related) is not: every call wakes the related object up and costs a
# query, which measured at ~110ms per row — hours over a full export. So the
# labels for a related table are fetched once, into {pk: label}, and looked up.
# Past LABEL_LIMIT rows that stops being worth it (and the big ones are things
# like Lead.duplicate_of pointing back at every lead), so the id stands alone.
LABEL_LIMIT = 3000


def _label_maps(cols, cache):
    """{field name: {pk: label}} for the foreign keys worth labelling."""
    maps = {}
    for _, f, kind in cols:
        if kind != 'label':
            continue
        target = f.related_model
        if target not in cache:
            cache[target] = ({o.pk: str(o) for o in target.objects.all()}
                             if target.objects.count() <= LABEL_LIMIT else None)
        maps[f.name] = cache[target]
    return maps


def _styled(ws, values, font, fill=None):
    """A row of write-only cells carrying their own styling."""
    from openpyxl.cell import WriteOnlyCell
    out = []
    for v in values:
        cell = WriteOnlyCell(ws, value=v)
        cell.font = font
        if fill:
            cell.fill = fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        out.append(cell)
    return out


def _append_table(ws, table, company, cache):
    cols = _columns(table.cls)
    maps = _label_maps(cols, cache)
    qs = table.cls.objects.filter(**{table.scope: company}).order_by('pk')
    count = qs.count()

    ws.append(_styled(ws, [table.label], Font(bold=True, size=12, color=NAVY)))
    ws.append(_styled(ws, [f'{count} row{"" if count == 1 else "s"}'],
                      Font(italic=True, size=9, color='FF6B7280')))
    ws.append(_styled(ws, [h for h, _, _ in cols],
                      Font(bold=True, color='FFFFFFFF', size=10),
                      PatternFill('solid', fgColor=NAVY)))

    # .iterator() so a table with six figures of rows never lands in memory at once.
    for obj in qs.iterator(chunk_size=2000):
        line = []
        for _, f, kind in cols:
            if kind == 'id':
                line.append(getattr(obj, f.attname, None))
            elif kind == 'label':
                m = maps.get(f.name)
                line.append(m.get(getattr(obj, f.attname, None)) if m else None)
            else:
                line.append(_cell(getattr(obj, f.name, None)))
        ws.append(line)

    ws.append([])   # blank spacer closes the table for the restore parser


def build_workbook(company):
    """A workbook holding everything this company owns, a sheet per module.

    write_only so the rows stream out instead of being held as cell objects —
    a full company runs to a couple of hundred thousand rows.
    """
    wb = openpyxl.Workbook(write_only=True)
    cache = {}
    for sheet_name, tables in SHEETS:
        ws = wb.create_sheet(sheet_name[:31])   # Excel caps sheet names at 31
        ws.append(_styled(ws, [f'{sheet_name} — {company.name}'],
                          Font(bold=True, size=14, color=NAVY)))
        ws.append(_styled(ws, [f'Backup taken {timezone.localtime().strftime("%d %b %Y, %I:%M %p")}'],
                          Font(italic=True, size=9, color='FF6B7280')))
        ws.append([])
        for table in tables:
            _append_table(ws, table, company, cache)
    return wb


# ── Restore ──────────────────────────────────────────────────────────────────
def parse_workbook(fileobj):
    """{table label: [row dicts]} for the restorable tables found in the file.

    Every table starts with its label alone in column A. Data rows always start
    with a numeric id, so a string there can only be a title or a header — which
    is what makes finding the sections reliable.
    """
    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    wanted = {t.label for t in RESTORABLE}
    found, label, headers = {}, None, None
    try:
        for ws in wb.worksheets:
            label, headers = None, None
            for row in ws.iter_rows(values_only=True):
                first = row[0] if row else None
                if isinstance(first, str) and first.strip() in ALL_LABELS:
                    label = first.strip()
                    headers = None
                    if label in wanted:
                        found.setdefault(label, [])
                    continue
                if label is None:
                    continue
                if not any(c is not None and c != '' for c in row):
                    label, headers = None, None     # blank row ends the table
                    continue
                if headers is None:
                    # The row count line sits between the title and the header.
                    if isinstance(first, str) and first.endswith(('row', 'rows')):
                        continue
                    headers = [str(c).strip() if c is not None else '' for c in row]
                    continue
                if label in wanted:
                    found[label].append(dict(zip(headers, row)))
    finally:
        wb.close()
    return found


def _empty_for(field):
    """What an empty cell means for this column.

    A blank-but-not-null column holds '' in the database, never NULL, and Excel
    gives back None for both — so handing None straight to a NOT NULL column is
    how a restore fails halfway. Fall back to the column's own idea of empty.
    """
    if field.null:
        return None
    if field.has_default():
        return field.get_default()
    if isinstance(field, (models.CharField, models.TextField)):
        return ''
    return None


def _to_python(field, value):
    """Turn a cell back into what the field expects."""
    if value is None or value == '':
        return _empty_for(field)
    if isinstance(field, models.DateTimeField):        # before DateField: it subclasses it
        if isinstance(value, datetime):
            dt = value                                  # an Excel date cell, from an older file
        elif isinstance(value, date):
            dt = datetime(value.year, value.month, value.day)
        else:
            from django.utils.dateparse import parse_datetime
            dt = parse_datetime(str(value).strip())
            if dt is None:
                return None
        # Exported ISO carries its offset, so this normally keeps the exact
        # original instant; a naive value (older file) is read as local time.
        return timezone.make_aware(dt) if timezone.is_naive(dt) else dt
    if isinstance(field, models.DateField):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        from django.utils.dateparse import parse_date
        return parse_date(str(value))
    if isinstance(field, models.TimeField):
        if isinstance(value, time):
            return value
        if isinstance(value, datetime):
            return value.time()
        from django.utils.dateparse import parse_time
        return parse_time(str(value))
    if isinstance(field, models.DecimalField):         # covers EncryptedDecimalField
        try:
            d = Decimal(str(value))
        except (InvalidOperation, ValueError, TypeError):
            return None
        # Excel hands back a bare number, so 2500000.00 arrives as 2500000 and
        # would be stored — and shown — a digit short. Put the scale back.
        if field.decimal_places is not None:
            d = d.quantize(Decimal(1).scaleb(-field.decimal_places))
        return d
    if isinstance(field, models.BooleanField):
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in ('1', 'true', 'yes', 'y')
    if isinstance(field, models.JSONField):
        if isinstance(value, (dict, list)):
            return value
        try:
            return json.loads(str(value))
        except (ValueError, TypeError):
            return {}
    if isinstance(field, (models.IntegerField, models.AutoField)):
        try:
            return int(value)
        except (ValueError, TypeError):
            return None
    if isinstance(field, models.FloatField):
        try:
            return float(value)
        except (ValueError, TypeError):
            return None
    if isinstance(field, models.UUIDField):
        try:
            return uuid.UUID(str(value))
        except (ValueError, TypeError, AttributeError):
            return None
    return str(value)


@contextmanager
def _keep_original_timestamps(model_classes):
    """Let auto_now / auto_now_add columns keep the values from the sheet.

    Django stamps those fields on write, which would silently replace every
    created_at with the moment of the restore — the one thing a restore must
    not do. Flipping the flags off is process-wide for the duration, so a
    concurrent write would also miss its stamp; acceptable here because this
    runs inside one short admin-only transaction.
    """
    touched = []
    for model in model_classes:
        for f in model._meta.fields:
            if getattr(f, 'auto_now', False) or getattr(f, 'auto_now_add', False):
                touched.append((f, f.auto_now, f.auto_now_add))
                f.auto_now = f.auto_now_add = False
    try:
        yield
    finally:
        for f, auto_now, auto_now_add in touched:
            f.auto_now, f.auto_now_add = auto_now, auto_now_add


def _as_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _foreign_rows(company, parsed):
    """Rows in the file that are not this company's.

    A workbook is restored into the company it was taken from. Without this the
    rows would be written with whatever company the *file* names, so one
    company's admin uploading another's backup would put that company's records
    into the database — the request being scoped to their own company would not
    stop it, because the scope is only used to choose the file, not the rows.
    """
    from django.apps import apps as _apps
    Lead = _apps.get_model('sales.Lead')
    User = _apps.get_model('accounts.User')

    lead_rows = parsed.get('Leads') or []
    mine = {_as_int(r.get('id')) for r in lead_rows
            if str(r.get('company_id') or '') == str(company.pk)}
    # Children may also point at leads this company already has (a partial reset).
    valid_leads = mine | set(Lead.objects.filter(company=company).values_list('pk', flat=True))
    valid_users = set(User.objects.filter(company=company).values_list('pk', flat=True))

    problems = []
    for table in RESTORABLE:
        rows = parsed.get(table.label) or []
        bad = 0
        for r in rows:
            if table.scope == 'company':
                own = str(r.get('company_id') or '') == str(company.pk)
                # Booking.company is nullable; fall back to the lead it belongs to.
                if not own and not r.get('company_id'):
                    own = _as_int(r.get('lead_id')) in valid_leads
            elif table.scope.startswith('lead__'):
                own = _as_int(r.get('lead_id')) in valid_leads
            elif table.scope.startswith('user__'):
                own = _as_int(r.get('user_id')) in valid_users
            elif table.scope.startswith('recipient__'):
                own = _as_int(r.get('recipient_id')) in valid_users
            else:
                own = True
            if not own:
                bad += 1
        if bad:
            problems.append({'table': table.label, 'rows': bad})
    return problems


def restore(company, parsed, commit=False):
    """Put the rows in `parsed` back, or report what would happen.

    Refuses outright if any row already exists: a restore is only ever meant to
    run into the empty slate Data Reset leaves behind, so an id that is already
    taken means this is not that situation and writing anything would be
    guesswork. Refuses too if the file holds another company's rows.
    """
    foreign = _foreign_rows(company, parsed)
    if foreign:
        plan = [{'table': t.label, 'rows': len(parsed.get(t.label) or [])} for t in RESTORABLE]
        return {'ok': False, 'committed': False, 'plan': plan,
                'total': sum(p['rows'] for p in plan), 'conflicts': [], 'foreign': foreign,
                'detail': f'This workbook holds records belonging to another company, so it '
                          f'cannot be restored into {company.name}. Nothing was written.'}

    plan, conflicts = [], []
    for table in RESTORABLE:
        rows = parsed.get(table.label) or []
        ids = [r.get('id') for r in rows if r.get('id') not in (None, '')]
        ids = [int(i) for i in ids]
        taken = list(table.cls.objects.filter(pk__in=ids).values_list('pk', flat=True)[:20]) if ids else []
        if taken:
            conflicts.append({'table': table.label, 'existing_ids': taken})
        plan.append({'table': table.label, 'rows': len(rows)})

    total = sum(p['rows'] for p in plan)
    if conflicts:
        return {'ok': False, 'committed': False, 'plan': plan, 'total': total,
                'conflicts': conflicts, 'foreign': [],
                'detail': 'Some of these records already exist. Restore only runs into '
                          'data that has been cleared — nothing was written.'}
    if not commit:
        return {'ok': True, 'committed': False, 'plan': plan, 'total': total,
                'conflicts': [], 'foreign': []}

    classes = [t.cls for t in RESTORABLE]
    with _keep_original_timestamps(classes), transaction.atomic():
        for table in RESTORABLE:
            rows = parsed.get(table.label) or []
            if not rows:
                continue
            model = table.cls
            by_header = {h: (f, kind) for h, f, kind in _columns(model) if kind != 'label'}
            objs = []
            for row in rows:
                kwargs = {}
                for header, raw in row.items():
                    spec = by_header.get(header)
                    if spec is None:
                        continue            # label column, or a column we don't write
                    f, kind = spec
                    if kind == 'id':
                        kwargs[f.attname] = _to_python(f.target_field, raw)
                    else:
                        kwargs[f.name] = _to_python(f, raw)
                objs.append(model(**kwargs))
            model.objects.bulk_create(objs, batch_size=500)

    return {'ok': True, 'committed': True, 'plan': plan, 'total': total,
            'conflicts': [], 'foreign': []}
