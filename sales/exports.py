"""Large list downloads (Leads, Site Visits) built in the background.

A company's whole lead list is tens of thousands of rows, every name, phone and
remark encrypted — decrypting and writing them takes longer than one web request
is allowed (gunicorn --timeout 60), so a single "build it and send it" request ran
out of time and the button sat on "Preparing…" for ever.

Instead the list view starts a job and answers at once with its id. The job runs in
a thread of the same worker and writes the workbook to a folder every worker of
this container shares; the screen polls its status (rows done / total) and fetches
the file when it is ready. Only the person who started a job can read or fetch it,
and a fetched or hour-old file is deleted.
"""
import json
import os
import threading
import time
import uuid

from django.db import connection
from django.utils import timezone

EXPORT_DIR = os.path.join('/tmp', 'nexora-exports')
MAX_AGE = 60 * 60          # seconds a finished file is kept if never fetched


def _paths(job):
    safe = ''.join(c for c in str(job) if c.isalnum())
    return os.path.join(EXPORT_DIR, safe + '.json'), os.path.join(EXPORT_DIR, safe + '.xlsx')


def _write_status(job, status):
    meta, _ = _paths(job)
    tmp = meta + '.tmp'
    with open(tmp, 'w') as f:
        json.dump(status, f)
    os.replace(tmp, meta)


def read_status(job):
    meta, _ = _paths(job)
    try:
        with open(meta) as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def file_path(job):
    return _paths(job)[1]


def remove(job):
    for p in _paths(job):
        try:
            os.remove(p)
        except OSError:
            pass


def _sweep():
    """Drop files left behind by jobs nobody came back for."""
    now = time.time()
    try:
        for name in os.listdir(EXPORT_DIR):
            p = os.path.join(EXPORT_DIR, name)
            if now - os.path.getmtime(p) > MAX_AGE:
                os.remove(p)
    except OSError:
        pass


def start(owner_id, total, rows_iter, headings, title, subtitle, filename):
    """Start building the workbook; returns the job id at once.

    rows_iter is a callable returning an iterable of row lists — called inside the
    thread, so the database work happens there and not in the request.
    """
    os.makedirs(EXPORT_DIR, exist_ok=True)
    _sweep()
    job = uuid.uuid4().hex
    status = {'owner': owner_id, 'status': 'running', 'done': 0, 'total': total,
              'filename': filename, 'started': timezone.now().isoformat()}
    _write_status(job, status)

    def run():
        try:
            import openpyxl
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.Workbook(write_only=True)
            ws = wb.create_sheet(title[:31])
            navy = 'FF0F1838'
            t = WriteOnlyCell(ws, value=title); t.font = Font(bold=True, size=14, color=navy)
            ws.append([t])
            st = WriteOnlyCell(ws, value=subtitle); st.font = Font(size=10, color='FF8492A6')
            ws.append([st])
            head = []
            for h in headings:
                c = WriteOnlyCell(ws, value=h)
                c.font = Font(bold=True, color='FFFFFFFF', size=10)
                c.fill = PatternFill('solid', fgColor=navy)
                head.append(c)
            ws.append(head)
            for i, h in enumerate(headings):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = min(40, max(12, len(h) + 4))
            ws.freeze_panes = 'A4'
            done = 0
            for row in rows_iter():
                ws.append(row)
                done += 1
                if done % 1000 == 0:
                    status['done'] = done
                    _write_status(job, status)
            _, xlsx = _paths(job)
            wb.save(xlsx)
            status.update(status='done', done=done)
            _write_status(job, status)
        except Exception as e:                       # the screen says what went wrong
            status.update(status='error', detail='Could not build the file: %s' % e)
            _write_status(job, status)
        finally:
            if not inline:
                connection.close()                   # this thread's own DB connection

    # Tests run it in line: a thread has its own connection and cannot see the
    # test's uncommitted rows.
    from django.conf import settings
    inline = getattr(settings, 'EXPORTS_INLINE', False)
    if inline:
        run()
    else:
        threading.Thread(target=run, daemon=True).start()
    return job
